"""The files connector: folders the person asked Alpha to read.

A watched folder is a connection. Syncing walks it, reads every new or changed document into
the world (text extracted; the file itself is never changed), makes each one a `document`
entity keyed by its path, and journals what it saw. A file that disappears is marked removed.
While the core runs, FSEvents (through watchdog) triggers the same sync for the folder.
"""

from __future__ import annotations

import hashlib
import json
import logging
import re
import shutil
import sqlite3
import threading
from collections.abc import Callable
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

from alpha.config import settings
from alpha.connectors.base import Connections
from alpha.world.store import Problem, fts_query, new_id, now
from alpha.world.world import World, alpha_home

log = logging.getLogger("alpha.files")

TEXT_SUFFIXES = {".md", ".markdown", ".txt", ".csv", ".tsv", ".json", ".yaml", ".yml", ".html",
                 ".htm", ".rtf", ".log"}
OFFICE_SUFFIXES = {".pdf", ".docx", ".xlsx", ".xls", ".pptx"}
SKIP_DIRS = {".git", "node_modules", ".venv", "__pycache__", ".Trash", "Library", ".cache"}
MAX_FILE_BYTES = settings.files.max_file_bytes
MAX_TEXT_CHARS = settings.files.max_text_chars


def column_letter(n: int) -> str:
    """1 → A, 27 → AA: a spreadsheet's own name for a column."""
    out = ""
    while n > 0:
        n, rest = divmod(n - 1, 26)
        out = chr(65 + rest) + out
    return out


def column_number(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n


def _value_text(value: Any) -> str:
    """A cell's value as it reads: a whole float without its .0, and a float's binary noise
    rounded away (Excel's 380881.41000000003 is 380881.41)."""
    if isinstance(value, float):
        value = round(value, 9)
        if value.is_integer():
            return str(int(value))
        return repr(value)
    return str(value)


def sheets_of(path: Path) -> list[dict[str, Any]]:
    """A spreadsheet's sheets cell by cell (P6, 19 Oct 2026): each sheet's name, its size, and
    its rows as {row: {column letter: (value, formula or None)}}. An .xlsx is read twice, for
    the values and for the formulas; an .xls through xlrd (values); a .csv or .tsv as one sheet.
    Empty cells are left out."""
    suffix = path.suffix.lower()
    sheets: list[dict[str, Any]] = []
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        values = load_workbook(str(path), read_only=True, data_only=True)
        formulas = load_workbook(str(path), read_only=True, data_only=False)
        for ws_v, ws_f in zip(values.worksheets, formulas.worksheets, strict=False):
            rows: dict[int, dict[str, tuple[Any, str | None]]] = {}
            width = 0
            for r, (row_v, row_f) in enumerate(zip(ws_v.iter_rows(min_row=1, values_only=True),
                                                    ws_f.iter_rows(min_row=1, values_only=True),
                                                    strict=False), start=1):
                cells: dict[str, tuple[Any, str | None]] = {}
                for c, (v, f) in enumerate(zip(row_v, row_f, strict=False), start=1):
                    formula = f if isinstance(f, str) and f.startswith("=") else None
                    if v is None and formula is None:
                        continue
                    cells[column_letter(c)] = (v, formula)
                    width = max(width, c)
                if cells:
                    rows[r] = cells
                if len(rows) > 20_000:
                    break
            sheets.append({"name": ws_v.title, "rows": rows,
                           "height": max(rows) if rows else 0, "width": width})
    elif suffix == ".xls":
        import xlrd  # type: ignore[import-untyped]

        book = xlrd.open_workbook(str(path))
        for sh in book.sheets():
            rows = {}
            for r in range(min(sh.nrows, 20_000)):
                cells = {column_letter(c + 1): (sh.cell_value(r, c), None)
                         for c in range(sh.ncols) if sh.cell_value(r, c) not in ("", None)}
                if cells:
                    rows[r + 1] = cells
            sheets.append({"name": sh.name, "rows": rows, "height": sh.nrows,
                           "width": sh.ncols})
    elif suffix in (".csv", ".tsv"):
        import csv

        with path.open(newline="", errors="replace") as fh:
            reader = csv.reader(fh, delimiter="\t" if suffix == ".tsv" else ",")
            rows, width = {}, 0
            for r, row in enumerate(reader, start=1):
                cells = {column_letter(c): (v, None) for c, v in enumerate(row, start=1)
                         if v.strip()}
                width = max(width, len(row))
                if cells:
                    rows[r] = cells
                if r > 20_000:
                    break
        sheets.append({"name": path.stem, "rows": rows, "height": max(rows) if rows else 0,
                       "width": width})
    return sheets


def _sheet_text(sheet: dict[str, Any]) -> str:
    lines = [f"# Sheet \"{sheet['name']}\" ({sheet['height']} rows × {sheet['width']} columns)"]
    for r, cells in sheet["rows"].items():
        lines.append(f"{r}: " + " | ".join(
            f"{col}={'' if v is None else _value_text(v)}" + (f" ({f})" if f else "")
            for col, (v, f) in cells.items()))
    return "\n".join(lines)


def read_document(path: Path) -> tuple[str, dict[str, Any]]:
    """A document's readable text with its positions, and its structure (P6): a PDF page by
    page ("[page 3]"); a sheet cell by cell ("12: A=Total revenue | B=48200 (=SUM(B2:B11))");
    the structure keeps the pages, or each sheet's name and size, so a value can be pinned to a
    place that exists (P5). Raises Problem with a plain reason when it can't."""
    suffix = path.suffix.lower()
    try:
        if suffix in (".xlsx", ".xls", ".csv", ".tsv"):
            sheets = sheets_of(path)
            structure = {"kind": "sheets", "sheets": [
                {"name": sh["name"], "rows": sh["height"], "columns": sh["width"]}
                for sh in sheets]}
            if suffix in (".csv", ".tsv"):
                return path.read_text(errors="replace")[:MAX_TEXT_CHARS], structure
            return "\n\n".join(_sheet_text(sh) for sh in sheets)[:MAX_TEXT_CHARS], structure
        if suffix in TEXT_SUFFIXES:
            return path.read_text(errors="replace")[:MAX_TEXT_CHARS], {"kind": "text"}
        if suffix == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            pages = [(p.extract_text() or "") for p in reader.pages]
            text = "\n\n".join(f"[page {i}]\n{t}" for i, t in enumerate(pages, start=1))
            return text[:MAX_TEXT_CHARS], {"kind": "pdf", "pages": len(pages)}
        if suffix == ".docx":
            import docx

            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)[:MAX_TEXT_CHARS], {"kind": "docx",
                                                         "paragraphs": len(d.paragraphs)}
        if suffix == ".pptx":
            from pptx import Presentation

            deck = Presentation(str(path))
            slides = []
            for i, slide in enumerate(deck.slides, 1):
                texts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
                slides.append(f"# Slide {i}\n" + "\n".join(texts))
            return "\n\n".join(slides)[:MAX_TEXT_CHARS], {"kind": "pptx",
                                                            "slides": len(slides)}
    except Exception as e:  # a damaged or protected file: say so, keep going
        raise Problem(f"Couldn't read {path.name}: {e}") from e
    raise Problem(f"Alpha doesn't read {suffix or 'this kind of'} files.")


def extract(path: Path) -> str:
    """Readable text of a document, with its positions. Raises Problem when it can't."""
    return read_document(path)[0]


def parse_range(text: str) -> tuple[int, int, int, int] | None:
    """ "B2:F40" → (first row, first column, last row, last column); "B12" is one cell."""
    m = re.fullmatch(r"\s*([A-Za-z]{1,3})(\d+)\s*(?::\s*([A-Za-z]{1,3})(\d+))?\s*", text)
    if not m:
        return None
    c1, r1 = column_number(m.group(1)), int(m.group(2))
    c2, r2 = (column_number(m.group(3)), int(m.group(4))) if m.group(3) else (c1, r1)
    return min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2)


def file_hash(path: Path) -> str:
    """The bytes' sha256: one file, one document (P6)."""
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def files_dir(module_name: str | None = None) -> Path:
    """Alpha's own folder for files it fetched or was given: one folder per module under the
    data directory (`files/<module>`), `files/loose` for the rest. Never the person's folders."""
    slug = "_".join("".join(ch if ch.isalnum() else " " for ch in (module_name or "loose")
                           .lower()).split()) or "loose"
    return alpha_home() / "files" / slug


def unique_path(folder: Path, name: str) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    clean = Path(name).name or "file"
    target = folder / clean
    n = 2
    while target.exists():
        target = folder / f"{Path(clean).stem} ({n}){Path(clean).suffix}"
        n += 1
    return target


def readable(path: Path) -> bool:
    return path.suffix.lower() in TEXT_SUFFIXES | OFFICE_SUFFIXES and not path.name.startswith(
        (".", "~$")
    )


def _walk(root: Path) -> list[Path]:
    out = []
    for p in root.rglob("*"):
        folders = p.relative_to(root).parts[:-1]
        if any(part in SKIP_DIRS or part.startswith(".") for part in folders):
            continue
        if p.is_file() and readable(p):
            out.append(p)
    return out


def _mtime(path: Path) -> str:
    return datetime.fromtimestamp(path.stat().st_mtime, UTC).replace(microsecond=0).isoformat()


def document_view(row: sqlite3.Row, *, full: bool = False) -> dict[str, Any]:
    out = {
        "id": row["id"],
        "entity_id": row["entity_id"],
        "path": row["path"],
        "title": row["title"],
        "kind": row["kind"],
        "size": row["size"],
        "modified_at": row["modified_at"],
        "chars": len(row["text"]),
        "module": row["module"] if "module" in row.keys() else None,
        "origin": row["origin"] if "origin" in row.keys() else None,
        # when Alpha took or first read it (slice 2 G6); a connected folder's own file
        "added_at": row["indexed_at"],
        "from_folder": bool(row["connection"]),
        "about": row["about"] if "about" in row.keys() else None,
        # where things are in it (P6): its pages, or its sheets and their sizes
        "structure": json.loads(row["structure"]) if "structure" in row.keys()
        and row["structure"] else None,
    }
    if full:
        out["text"] = row["text"]
    return out


class Files:
    def __init__(self, world: World) -> None:
        self.world = world
        self.connections = Connections(world.store)

    # ---- connections ----

    def watch(self, folder: str) -> dict[str, Any]:
        path = Path(folder).expanduser().resolve()
        if not path.is_dir():
            raise Problem(f"{folder} isn't a folder on this Mac.")
        if path == Path.home() or str(path) == "/":
            raise Problem("That's a whole drive or home folder; pick the folder that matters.")
        conn = self.connections.upsert("files", str(path))
        self.world.journal.append("made", f"Started reading the folder {path.name} ({path}).",
                                  data={"connection": conn["id"]}, source="connector:files")
        return conn

    # ---- syncing ----

    def sync(self, folder: str | None = None) -> dict[str, Any]:
        """Read new and changed documents in one watched folder (or all of them)."""
        targets = [c for c in self.connections.all("files") if c["status"] != "off"]
        if folder is not None:
            wanted = str(Path(folder).expanduser().resolve())
            targets = [c for c in targets if c["target"] == wanted]
            if not targets:
                raise Problem(f"Alpha isn't reading {folder}; ask it to watch the folder first.")
        totals = {"added": 0, "changed": 0, "removed": 0, "unreadable": 0, "folders": 0}
        for conn in targets:
            try:
                result = self._sync_one(conn)
            except Exception as e:
                log.exception("sync of %s failed", conn["target"])
                self.connections.synced(conn["id"], error=str(e))
                continue
            self.connections.synced(conn["id"])
            totals["folders"] += 1
            for k in ("added", "changed", "removed", "unreadable"):
                totals[k] += result[k]
        return totals

    def _sync_one(self, conn: dict[str, Any]) -> dict[str, int]:
        root = Path(conn["target"])
        if not root.is_dir():
            raise Problem(f"The folder {root} is gone.")
        store = self.world.store
        known = {
            r["path"]: r
            for r in store.all(
                "SELECT id, path, modified_at, removed_at FROM documents WHERE connection = ?",
                (conn["id"],),
            )
        }
        counts = {"added": 0, "changed": 0, "removed": 0, "unreadable": 0}
        seen: set[str] = set()
        for path in _walk(root):
            key = str(path)
            seen.add(key)
            if path.stat().st_size > MAX_FILE_BYTES:
                continue
            modified = _mtime(path)
            prior = known.get(key)
            if prior and prior["modified_at"] == modified and prior["removed_at"] is None:
                continue
            try:
                text, structure = read_document(path)
            except Problem as e:
                counts["unreadable"] += 1
                log.info("%s", e)
                continue
            self._save(conn, path, text, modified, prior, structure)
            counts["changed" if prior else "added"] += 1
        for path_str, prior in known.items():
            if path_str not in seen and prior["removed_at"] is None:
                with store.tx() as db:
                    db.execute("UPDATE documents SET removed_at = ? WHERE id = ?",
                               (now(), prior["id"]))
                self.world.journal.append(
                    "saw", f"{Path(path_str).name} was removed from {root.name}.",
                    data={"document": prior["id"]}, source="connector:files",
                )
                counts["removed"] += 1
        return counts

    def _save(self, conn: dict[str, Any], path: Path, text: str, modified: str,
              prior: sqlite3.Row | None, structure: dict[str, Any] | None = None) -> None:
        entity = self.world.entities.resolve("document", path.name, {"path": str(path)})["entity"]
        stamp = now()
        size = path.stat().st_size
        digest, shape = file_hash(path), json.dumps(structure) if structure else None
        with self.world.store.tx() as db:
            if prior:
                db.execute(
                    "UPDATE documents SET title = ?, text = ?, size = ?, modified_at = ?,"
                    " indexed_at = ?, removed_at = NULL, hash = ?, structure = ? WHERE id = ?",
                    (path.name, text, size, modified, stamp, digest, shape, prior["id"]),
                )
                did = prior["id"]
            else:
                did = new_id("d")
                db.execute(
                    "INSERT INTO documents (id, entity_id, connection, path, title, kind, text,"
                    " size, modified_at, indexed_at, hash, structure)"
                    " VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                    (did, entity["id"], conn["id"], str(path), path.name,
                     path.suffix.lower().lstrip("."), text, size, modified, stamp, digest,
                     shape),
                )
        folder = Path(conn["target"]).name
        words = len(text.split())
        self.world.journal.append(
            "saw",
            f"{'Read again' if prior else 'Read'} {path.name} in {folder} ({words:,} words).",
            data={"document": did, "path": str(path)}, entity_ids=[entity["id"]],
            source="connector:files",
        )

    # ---- taking a file in: fetched from a connection, or given by the person ----

    def take(self, source: Path, *, module: str | None, origin: str, name: str | None = None,
             move: bool = False, turn: str | None = None, by: str = "alpha",
             about: str | None = None, twin_ok: bool = False) -> dict[str, Any]:
        """Keep a file in Alpha's folder for the module and make it a document (its text
        extracted where Alpha reads the kind; otherwise kept as a file alone). `origin` says
        where it came from (a page or message address, or "the person"). Returns the
        document."""
        source = Path(source)
        if not source.is_file():
            raise Problem(f"There is no file at {source}.")
        digest = file_hash(source)
        kept = None if twin_ok else self.kept(digest)
        if kept is not None:
            # one file, one document (P6): the same bytes again are the document already kept
            if move:
                source.unlink(missing_ok=True)
            self.world.journal.append(
                "did", f"Already kept {kept['title']}: the same file from {origin}.",
                actor=by, data={"document": kept["id"], "origin": origin, "turn": turn,
                                "already": True}, source="connector:files")
            return {**kept, "already": True}
        module_row = self.world.modules.get(module) if module else None
        target = unique_path(files_dir(module_row["name"] if module_row else None),
                             name or source.name)
        if move:
            shutil.move(str(source), target)
        else:
            shutil.copy2(source, target)
        structure: dict[str, Any] | None = None
        try:
            text, structure = read_document(target) if readable(target) else ("", None)
        except Problem as e:
            log.info("%s", e)
            text = ""
        entity = self.world.entities.resolve("document", target.name,
                                             {"path": str(target)})["entity"]
        stamp = now()
        did = new_id("d")
        with self.world.store.tx() as db:
            db.execute(
                "INSERT INTO documents (id, entity_id, connection, path, title, kind, text, size,"
                " modified_at, indexed_at, module, origin, about, hash, structure)"
                " VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (did, entity["id"], None, str(target), target.name,
                 target.suffix.lower().lstrip("."), text, target.stat().st_size,
                 _mtime(target), stamp, module_row["id"] if module_row else None, origin,
                 about, digest, json.dumps(structure) if structure else None),
            )
        words = len(text.split())
        person = self.world.entities.get(about) if about else None
        where = (f" into {module_row['name']}" if module_row else
                 f" about {person['name']}" if person else "")
        self.world.journal.append(
            "did" if by == "alpha" else "changed",
            (f"Kept {target.name}{where} from {origin}" if by == "alpha" else
             f"Added {target.name}{where}")
            + (f" ({words:,} words read)" if words else " (kept as a file; not a kind Alpha"
                                                        " reads)") + ".",
            actor=by, data={"document": did, "path": str(target), "origin": origin,
                            "turn": turn},
            entity_ids=[entity["id"]] + ([about] if about else []),
            module=module_row["id"] if module_row else None,
            source=None if by == "person" else "connector:files",
        )
        return self.document(did)

    def kept(self, digest: str) -> dict[str, Any] | None:
        """The document Alpha keeps with these bytes, if any (a connected folder's own files
        are the person's and never count: they are read where they are)."""
        row = self.world.store.one(
            "SELECT * FROM documents WHERE hash = ? AND removed_at IS NULL AND connection IS NULL"
            " ORDER BY indexed_at LIMIT 1", (digest,))
        return document_view(row) if row else None

    def document(self, did: str) -> dict[str, Any]:
        row = self.world.store.one("SELECT * FROM documents WHERE id = ?", (did,))
        if row is None:
            raise Problem(f"There is no document {did}.")
        return document_view(row)

    def of_module(self, module_id: str) -> list[dict[str, Any]]:
        return [document_view(r) for r in self.world.store.all(
            "SELECT * FROM documents WHERE module = ? AND removed_at IS NULL"
            " ORDER BY indexed_at DESC", (module_id,))]

    def of_entity(self, entity_id: str) -> list[dict[str, Any]]:
        """The files kept about a person or company (B11), newest first."""
        return [document_view(r) for r in self.world.store.all(
            "SELECT * FROM documents WHERE about = ? AND removed_at IS NULL"
            " ORDER BY indexed_at DESC", (entity_id,))]

    def used_by(self, did: str) -> list[dict[str, Any]]:
        """The records whose file fields hold this file: table, record, its title, the field."""
        out: list[dict[str, Any]] = []
        for row in self.world.store.all(
            "SELECT collection, id, \"values\" FROM records WHERE deleted_at IS NULL AND"
            " \"values\" LIKE ?", (f"%{did}%",)):
            desc = self.world.collections.describe(row["collection"])
            values = json.loads(row["values"] or "{}")
            for f in desc["fields"]:
                if f["kind"] == "file" and values.get(f["name"]) == did:
                    title = values.get(desc.get("title_field") or "") or row["id"]
                    out.append({"collection": row["collection"], "table": desc["title"],
                                "record": row["id"], "title": str(title), "field": f["name"]})
        return out

    # ---- reading ----

    def search(self, text: str, limit: int = 10) -> list[dict[str, Any]]:
        query = fts_query(text)
        if query is None:
            return []
        rows = self.world.store.all(
            "SELECT d.*, snippet(documents_fts, 1, '[', ']', '…', 14) AS snip FROM documents_fts"
            " JOIN documents d ON d.rowid = documents_fts.rowid WHERE documents_fts MATCH ?"
            " AND d.removed_at IS NULL ORDER BY bm25(documents_fts) LIMIT ?",
            (query, limit),
        )
        return [{**document_view(r), "snippet": r["snip"]} for r in rows]

    def read(self, ref: str, start: int = 0, length: int = 20_000, *, page: int | None = None,
             sheet: str | None = None, cells: str | None = None) -> dict[str, Any]:
        """A document's text from `start`; or one PDF page; or a sheet's cells in a range
        ("B2:F40", the whole sheet when none), read again from the kept file (P6)."""
        row = self.world.store.one(
            "SELECT * FROM documents WHERE (id = ? OR path = ? OR title = ?)"
            " AND removed_at IS NULL",
            (ref, ref, ref),
        )
        if row is None:
            raise Problem(f"Alpha hasn't read a document '{ref}'.")
        view = document_view(row, full=True)
        if page is not None:
            pages = (view.get("structure") or {}).get("pages")
            if not pages:
                raise Problem(f"{view['title']} has no pages to read by number.")
            if not 1 <= page <= pages:
                raise Problem(f"{view['title']} has pages 1 to {pages}.")
            marker = f"[page {page}]\n"
            at = view["text"].find(marker)
            end = view["text"].find(f"[page {page + 1}]\n")
            view["text"] = view["text"][at + len(marker): end if end >= 0 else None] \
                if at >= 0 else ""
            view["page"] = page
            view["more"] = page < pages
            return view
        if sheet is not None or cells is not None:
            found = self.cells(view, sheet, cells)
            view["text"] = "\n".join(f"{c['cell']} = {c['value']}"
                                     + (f" ({c['formula']})" if c["formula"] else "")
                                     for c in found["cells"])
            view.update(sheet=found["sheet"], cells=found["cells"], more=False)
            return view
        view["text"] = view["text"][start : start + length]
        view["more"] = start + length < view["chars"]
        return view

    def sheet_table(self, doc: dict[str, Any], sheet: str | None,
                    header_row: int | None = None) -> dict[str, Any]:
        """A sheet read as a table (P6's import): its header row (the one given, else the
        first row of two or more cells that are all words) and the rows below it, each cell
        with its reference. Read again from the kept file."""
        sheets = sheets_of(Path(doc["path"]))
        chosen = next((sh for sh in sheets if sheet is None or sh["name"] == sheet), None)
        if chosen is None:
            raise Problem(f"{doc['title']} has no sheet {sheet!r}; its sheets are"
                          f" {', '.join(sh['name'] for sh in sheets) or 'none'}.")
        rows = chosen["rows"]

        def filled(cells: dict[str, Any]) -> list[Any]:
            return [v for v, _ in cells.values() if v not in (None, "")]

        if header_row is None:
            # an export's title rows above it hold one cell each; its corner may be empty
            # (a QuickBooks report's account column has no name)
            header_row = next((r for r, cells in rows.items() if len(filled(cells)) >= 2 and all(
                isinstance(v, str) and v.strip() for v in filled(cells))), None)
        if header_row is None or header_row not in rows:
            raise Problem(f"Couldn't find the header row of {chosen['name']}; say which row"
                          " holds the column names (header_row).")
        # a column with no name is called by its letter, so `columns` can name it ("A")
        headers = {col: str(v).strip() if v not in (None, "") else col
                   for col, (v, _) in rows[header_row].items()}
        body = [(r, cells) for r, cells in rows.items() if r > header_row]
        for _, cells in body:
            for col in cells:
                headers.setdefault(col, col)
        headers = dict(sorted(headers.items(), key=lambda kv: column_number(kv[0])))
        return {"sheet": chosen["name"], "header_row": header_row, "headers": headers,
                "rows": body}

    def cells(self, doc: dict[str, Any], sheet: str | None,
              cells: str | None = None) -> dict[str, Any]:
        """A sheet's cells, read again from the kept file: each with its reference, value and
        formula. A sheet that isn't there, or a range outside it, is refused in plain words."""
        sheets = sheets_of(Path(doc["path"]))
        if not sheets:
            raise Problem(f"{doc['title']} has no sheets.")
        names = [sh["name"] for sh in sheets]
        chosen = next((sh for sh in sheets if sheet is None or sh["name"] == sheet), None)
        if chosen is None:
            raise Problem(f"{doc['title']} has no sheet {sheet!r}; its sheets are"
                          f" {', '.join(names)}.")
        lo_row, hi_row, lo_col, hi_col = 1, chosen["height"], 1, chosen["width"]
        if cells:
            box = parse_range(cells)
            if box is None:
                raise Problem(f"{cells!r} isn't a range like B2:F40 or B12.")
            lo_row, lo_col, hi_row, hi_col = box
        out = [{"cell": f"{col}{r}", "value": "" if v is None else _value_text(v), "formula": f}
               for r, row in chosen["rows"].items() if lo_row <= r <= hi_row
               for col, (v, f) in row.items() if lo_col <= column_number(col) <= hi_col]
        return {"sheet": chosen["name"], "cells": out}

    def place(self, ref: str, *, page: int | None = None, sheet: str | None = None,
              cell: str | None = None) -> dict[str, Any]:
        """Where a value sat (P5), for the person to see: a PDF's page with its text, or a
        sheet's cell among its neighbours (three rows and two columns either side), each read
        again from the kept file."""
        doc = self.read(ref, 0, 1)
        out: dict[str, Any] = {"document": doc["id"], "title": doc["title"],
                               "path": doc["path"]}
        if page is not None:
            return {**out, "page": page, "text": self.read(doc["id"], page=page)["text"]}
        box = parse_range(cell) if cell else None
        if box is None or cell is None:
            return out
        top, left, bottom, right = box
        sheets = (doc.get("structure") or {}).get("sheets") or []
        size = next((sh for sh in sheets if sheet is None or sh["name"] == sheet), None)
        height = int(size["rows"]) if size else bottom + 3
        width = int(size["columns"]) if size else right + 2
        r1, r2 = max(1, top - 3), min(max(height, bottom), bottom + 3)
        c1, c2 = max(1, left - 2), min(max(width, right), right + 2)
        found = self.cells(doc, sheet, f"{column_letter(c1)}{r1}:{column_letter(c2)}{r2}")
        held = {c["cell"]: c for c in found["cells"]}
        columns = [column_letter(c) for c in range(c1, c2 + 1)]
        rows = [{"row": r, "cells": [{"cell": f"{col}{r}", "value": held.get(f"{col}{r}", {})
                                      .get("value", ""), "formula": held.get(f"{col}{r}", {})
                                      .get("formula")} for col in columns]}
                for r in range(r1, r2 + 1)]
        return {**out, "sheet": found["sheet"], "cell": cell.replace(" ", "").upper(),
                "columns": columns, "rows": rows}

    def documents(self, limit: int = 100) -> list[dict[str, Any]]:
        rows = self.world.store.all(
            "SELECT * FROM documents WHERE removed_at IS NULL ORDER BY modified_at DESC LIMIT ?",
            (limit,),
        )
        return [document_view(r) for r in rows]

    # ---- live updates while the core runs ----

    def observe(
        self, on_change: Callable[[dict[str, Any]], None] | None = None
    ) -> Callable[[], None]:
        """Watch every connected folder; a burst of changes syncs that folder once. Returns a
        function that stops watching."""
        from watchdog.events import FileSystemEvent, FileSystemEventHandler
        from watchdog.observers import Observer

        timers: dict[str, threading.Timer] = {}
        lock = threading.Lock()

        def run(folder: str) -> None:
            try:
                result = self.sync(folder)
                if on_change and any(result[k] for k in ("added", "changed", "removed")):
                    on_change({"folder": folder, **result})
            except Exception:
                log.exception("live sync of %s failed", folder)

        class Handler(FileSystemEventHandler):
            def __init__(self, folder: str) -> None:
                self.folder = folder

            def on_any_event(self, event: FileSystemEvent) -> None:
                if event.is_directory:
                    return
                with lock:
                    old = timers.pop(self.folder, None)
                    if old:
                        old.cancel()
                    timer = threading.Timer(2.0, run, args=(self.folder,))
                    timer.daemon = True
                    timers[self.folder] = timer
                    timer.start()

        observer = Observer()
        for conn in self.connections.all("files"):
            if conn["status"] != "off" and Path(conn["target"]).is_dir():
                observer.schedule(Handler(conn["target"]), conn["target"], recursive=True)
        observer.daemon = True
        observer.start()

        def stop() -> None:
            observer.stop()
            with lock:
                for t in timers.values():
                    t.cancel()

        return stop
