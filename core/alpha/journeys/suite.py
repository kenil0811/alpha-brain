"""Run the journeys in `journeys/*.yaml` against a copy of a world and write a report.

A journey file:

    name: branded_food
    title: A branded food is logged from its label, not guessed
    steps:
      - say: "i had a 45g bar of cadbury dairy milk"      # a turn, as the person
        module: Nutrition
      - reader: {name: linkedin_connections, into: linkedin_connections, key: profile_url}
      - agent: "Daily deal tracker"                        # run one agent now, by name
      - approve_action: latest   # (in steps_after) the person's yes to the newest action
      - seed: {kind: turn, ago: 1d, said: "…", replied: "…"}   # a past exchange
      - seed: {kind: fact, predicate: height_cm, value: "178"}      # a known fact
      - remove_module: "Car Wash Site Scoring"   # the world without something it already has
    checks:
      - independent: {}                     # the second opinion on the last turn agrees
      - row: {collection: food_log, source_not: [estimated]}   # a row this journey added
      - near: {collection: food_log, field: calories, value: 240, within: 0.1}
      - no_new_tables: {}
      - no_new_modules: {}
      - workflow: {state: drafted, steps_min: 2}
      - reply: {matches: "\\n\\s*1[.)]"}     # regex on the last reply (or `contains`)
      - judge: "Did the answer name people from the connections table, or say plainly none?"
      - judge: {rubric: "…", history: 1}   # the judge also sees the previous exchange
      - count: {collection: linkedin_connections, at_least_fraction_of_last_ok: 0.95,
                reader: linkedin_connections}
      - reader_health: linkedin_connections
      - agent: {title: "Daily deal tracker", result_matches: "Read \\d+ of \\d+",
                     problem_allowed: "sign-in|automated reading"}
      - journal: {kind: saw, url_contains: mail.google.com}   # an entry since the journey began

Every step and check is timed; every verdict says why. The world the suite runs on is a copy
(SQLite's backup, plus the browser profiles so signed-in sites read as the person), made in a
scratch directory the report names; the person's own world is never touched.
"""

from __future__ import annotations

import json
import logging
import os
import re
import shutil
import sqlite3
import threading
import time
from dataclasses import dataclass, field
from datetime import UTC, datetime, timedelta
from datetime import datetime as dt
from pathlib import Path
from typing import Any

import yaml
from fastapi.testclient import TestClient

from alpha import config
from alpha.api.server import create_app
from alpha.context import layers
from alpha.mcp.tools import Tools
from alpha.runtime import (
    acting,
    conversations,
    desk,
    engine,
    look,
    pipeline,
    route,
    running,
    staging,
    turn,
)
from alpha.runtime.agents import builtin, columbo, look_and_run
from alpha.runtime.agents import runner as agent_runner
from alpha.world import grants, purge
from alpha.world import people as people_world
from alpha.world.store import Problem, new_id, now
from alpha.world.tasks import TABLE as TASKS
from alpha.world.world import World, default_world_path

log = logging.getLogger(__name__)

APP_HOME = Path.home() / "Library" / "Application Support" / "com.alpha.brain"


def journeys_dir() -> Path:
    configured = os.environ.get("ALPHA_JOURNEYS")
    if configured:
        return Path(configured)
    return Path(__file__).resolve().parents[3] / "journeys"


def load(names: list[str] | None = None, folder: Path | None = None) -> list[dict[str, Any]]:
    folder = folder or journeys_dir()
    out = []
    for path in sorted(folder.glob("*.yaml")):
        try:
            data = yaml.safe_load(path.read_text()) or {}
        except yaml.YAMLError as e:
            raise Problem(f"The journey file {path.name} isn't valid YAML: {e}") from e
        data.setdefault("name", path.stem)
        if names and data["name"] not in names:
            continue
        out.append(data)
    if names:
        missing = set(names) - {j["name"] for j in out}
        if missing:
            raise Problem(f"No journey named {', '.join(sorted(missing))} in {folder}.")
    return out


# ---- the copy ----


def _titled(sg: dict[str, Any]) -> str:
    """A suggestion's title, with who brought it when an agent did."""
    return str(sg["title"]) + (f" [{sg.get('by_name')}]" if sg.get("by") else "")


def fresh_home(source_world: Path, into: Path) -> Path:
    """A scratch ALPHA_HOME with a new, empty world (the first-day sentences of the harness
    design §17 start from nothing) and the app's browser profiles, so a sign-in the person made
    holds. Returns the new world's path."""
    into.mkdir(parents=True, exist_ok=True)
    target = into / "world.sqlite"
    _copy_profiles(source_world, into)

    return target


def copy_home(source_world: Path, into: Path) -> Path:
    """A scratch ALPHA_HOME holding a consistent copy of the world and the browser profiles.
    Returns the copied world's path."""
    into.mkdir(parents=True, exist_ok=True)
    target = into / "world.sqlite"
    src = sqlite3.connect(f"file:{source_world}?mode=ro", uri=True)
    dst = sqlite3.connect(target)
    try:
        src.backup(dst)
    finally:
        dst.close()
        src.close()
    _copy_profiles(source_world, into)
    # the skills and the modules' files live beside the world (2b, slice 5): a copy without
    # them has agents whose run skills are missing (the check of 18 Oct found the page failing)
    for name in ("skills", "modules"):
        beside = source_world.parent / name
        if beside.is_dir():
            shutil.copytree(beside, into / name, dirs_exist_ok=True,
                            ignore=shutil.ignore_patterns("documents", "inputs", "parts"))
    return target


def _copy_profiles(source_world: Path, into: Path) -> None:
    profiles = source_world.parent / "browser"
    if profiles.is_dir():
        try:
            shutil.copytree(profiles, into / "browser", dirs_exist_ok=True,
                            ignore_dangling_symlinks=True,
                            ignore=shutil.ignore_patterns("Singleton*", "lockfile", "*.lock",
                                                          "RunningChromeVersion"))
        except shutil.Error as e:  # a file Chrome had open vanished mid-copy: not a cookie
            log.warning("profile copy skipped %d files: %s", len(e.args[0]), e.args[0][:2])
    (into / "logs").mkdir(exist_ok=True)


# ---- running ----


@dataclass
class Mark:
    """Where a journey began: what existed before it, to judge what it added."""
    at: str
    tables: set[str]
    modules: set[str]
    skills: dict[str, int]


def mark(world: World) -> Mark:
    return Mark(
        at=now(),  # UTC, like every `at` and `created_at` in the world
        tables={t["name"] for t in world.collections.overview()},
        modules={m["id"] for m in world.modules.all()},
        skills={s["name"]: s["version"] for s in world.skills.all()},
    )


@dataclass
class Outcome:
    name: str
    title: str
    steps: list[dict[str, Any]] = field(default_factory=list)
    checks: list[dict[str, Any]] = field(default_factory=list)
    seconds: float = 0.0
    error: str | None = None

    @property
    def passed(self) -> bool:
        return self.error is None and all(c["ok"] for c in self.checks)


class Run:
    def __init__(self, world: World, runner: turn.Runner = route.run) -> None:
        self.world = world
        self.runner = runner
        self.last_turn: str | None = None
        self.last_reply: str = ""
        self.exchanges: list[tuple[str, str]] = []  # (said, reply) in order, this journey
        self.last_agent: dict[str, Any] | None = None
        self.last_action: dict[str, Any] | None = None
        self.last_row: tuple[str, str] | None = None
        self.last_kit: str | None = None
        self.last_document: dict[str, Any] | None = None
        self.last_thread: str | None = None  # the conversation the last sentence went to
        # hosts allowed on their cards during a turn, with the card's module
        self.allowed_hosts: list[tuple[str, str | None]] = []
        builtin.ensure(world)
        self.mark = mark(world)

    def _desk_runner(self) -> turn.Runner:
        """The desk of this journey's world when the journey thinks with the real model, as
        the app's sentences do (section C); a fake runner in the suite's own tests."""
        if self.runner is route.run:
            return desk.of(self.world).run
        return self.runner

    # steps

    def step(self, spec: dict[str, Any]) -> dict[str, Any]:
        began = time.monotonic()
        if "say" in spec:
            # A hand that asks for the person's yes on a card (a row's removal, harness §8)
            # gets it from the suite: on a scratch world, yes, journaled like any answer.
            stop_answering = threading.Event()
            # `answer_cards: false` leaves the yes cards of work beside the desk to the
            # person's words (C4); one the desk's own turn waits on is still answered as their
            # button press would be, since the desk takes no sentence until it is (Q108)
            threading.Thread(target=self._answer_permissions, args=(stop_answering,),
                             kwargs={"desk_only": not spec.get("answer_cards", True)},
                             daemon=True).start()
            try:
                # Every sentence is the desk's (section C, Q107): Home's session, kept open,
                # with no routing before it, as the app does. `home:` and `here:` once chose
                # between routing and an item's own box; both are the desk now.
                home = self.world.modules.home_thread()["id"]
                runner = self._desk_runner()
                outcome = turn.ask(self.world, spec["say"], module=spec.get("module"),
                                   thread=home, runner=runner, conversation=True)
                deferred = [a for a in self.world.journal.open_asks()
                            if a["data"].get("tool_use")
                            and a["data"].get("turn") == outcome.said]
                if deferred:
                    # a removal ended the turn on its card (P3 item 7): the suite answers yes
                    # as the app does, and the desk's next turn carries on
                    self.world.journal.append(
                        "answered", "Yes, remove it", actor="person",
                        data={"ask": deferred[-1]["id"]}, thread=home)
                    outcome = turn.ask(self.world, "Yes, remove it", thread=home,
                                       runner=runner, conversation=True)
                while self.allowed_hosts:
                    # hosts the person allowed during the turn: the desk is told, and runs
                    # again what needed them (Q109)
                    pairs, self.allowed_hosts = list(dict.fromkeys(self.allowed_hosts)), []
                    hosts = [h for h, _ in pairs]
                    outcome = turn.ask(self.world,
                                       " ".join(engine.ALLOWED.format(host=h) for h in hosts),
                                       module=pairs[0][1], thread=home, runner=runner,
                                       conversation=True, actor="alpha",
                                       journal_as=f"{', '.join(hosts)} now allowed for scripts.")
                conversations.settle(self.world, home, ok=outcome.ok,
                                     stopped=outcome.result.stopped,
                                     last_line=conversations.first_line(outcome.reply))
                self.last_thread = home
            finally:
                stop_answering.set()
            self.last_turn, self.last_reply = outcome.said, outcome.reply
            self.exchanges.append((spec["say"], outcome.reply))
            if outcome.ok:
                # What the app's kick does after a person's turn: Watson looks and runs, in line.
                try:
                    look_and_run(self.world, "watson", runner=self.runner)
                except Exception as e:  # a journey is not failed by Watson's pass
                    log.warning("watson in journey: %s", e)
            record = {"say": spec["say"], "ok": outcome.ok, "reply": outcome.reply[:600],
                      "steps": outcome.result.num_turns}
            # what the turn read in tokens and when its first words came (B8 is measured by
            # it: the journey's world is thrown away, its report stays)
            replied = self.world.journal.read(outcome.replied)["data"]
            record.update({k: replied[k] for k in (*turn.TOKENS, "first_text_ms")
                           if replied.get(k) is not None})
        elif "reader" in spec:
            r = spec["reader"]
            out = pipeline.run_reader(self.world, r["name"], r["into"], r["key"],
                                      keep=r.get("keep"))
            record = {"reader": r["name"], "ok": out.get("health") == "ok",
                      "health": out.get("health"), "detail": out.get("problem") or out.get("line")}
        elif "agent" in spec:
            agent = self._agent(spec["agent"])
            self.last_agent = agent_runner.run(self.world, agent["id"], trigger="now",
                                               runner=self.runner)
            record = {"agent": agent["name"], "ok": not self.last_agent.get("last_error"),
                      "result": self.last_agent.get("last_result"),
                      "problem": self.last_agent.get("last_error")}
        elif "wait_for_card" in spec:
            # work beside the desk asks the person's yes on a card (a host its script reaches)
            # while the desk is free: the step waits until `count` such cards are open
            want = int((spec["wait_for_card"] or {}).get("count", 1))
            end = time.monotonic() + float((spec["wait_for_card"] or {}).get("seconds", 300))
            open_now: list[dict[str, Any]] = []
            while time.monotonic() < end:
                open_now = [a for a in self.world.journal.open_asks()
                            if a["data"].get("permission")]
                if len(open_now) >= want:
                    break
                time.sleep(1)
            record = {"wait_for_card": want, "ok": len(open_now) >= want,
                      "detail": "; ".join(a["text"][:80] for a in open_now) or "none open"}
        elif "wait_for_work" in spec:
            # what the person does while work runs beside the desk (section C): nothing; the
            # step waits until every piece of work this journey started has ended and the
            # desk has said what came of it
            limit = float((spec["wait_for_work"] or {}).get("seconds", 900))
            end = time.monotonic() + limit
            unfinished: list[str] = []
            while time.monotonic() < end:
                unfinished = [str(w["title"]) for w in self._work()
                              if w["state"] in ("working", "needs_you")
                              or not self._outcome_reply(str(w["id"]))]
                if not unfinished:
                    break
                time.sleep(2)
            # the desk's words on the work are the answer the checks read from now on
            spoken = [r["text"] for w in self._work()
                      if (r := self._outcome_reply(str(w["id"])))]
            if spoken:
                self.last_reply = "\n\n".join(dict.fromkeys(spoken))
            record = {"wait_for_work": len(self._work()), "ok": not unfinished,
                      "detail": ("all ended, each with the desk's words on it" if not unfinished
                                 else "still waiting on: " + "; ".join(unfinished))}
        elif "compact" in spec:
            # the last sentence's conversation is compacted as a long one would be (B6): the
            # SDK's own /compact, in that session; it passes when the summary ran and the start
            # layer was given again after it
            thread = self.last_thread or self.world.modules.home_thread()["id"]
            squeezed = turn.ask(self.world, "/compact", thread=thread,
                                runner=self._desk_runner(), conversation=True,
                                journal_as="(the journey compacts this conversation)")
            kept = self.world.journal.context(squeezed.said)
            again = bool(kept and "GIVEN AGAIN AFTER A COMPACTION" in str(kept["context"]))
            record = {"compact": self.world.modules.thread(thread)["title"], "ok": again,
                      "detail": "compacted; the start layer given again" if again
                      else "no compaction came back"}
        elif "start_tasks" in spec:
            # what the scheduler's tick does (P2): the tasks agents were given and someone moved
            # to Working, run now with the journey's model
            started = []
            for task in self.world.tasks.waiting_to_start():
                agent_runner.run(self.world, str(task["owner"]), trigger="task",
                                 task=task["id"], runner=self.runner)
                started.append(task["title"])
            record = {"start_tasks": started, "ok": bool(started),
                      "detail": f"{len(started)} task(s) started" + (
                          f": {'; '.join(started)}" if started else "")}
        elif "file_document" in spec:
            # a file the person dropped into a module (P6): a sheet or a PDF built from the
            # journey's own words, kept as Alpha keeps any file; `again` brings the same bytes
            # under another name, as a second email would
            arg = spec["file_document"]
            doc = self._file_document(arg)
            self.last_document = doc
            record = {"file_document": doc["title"], "ok": True,
                      "detail": ("already kept" if doc.get("already") else "kept")
                      + f" as {doc['id']}"}
        elif "import_document" in spec:
            # the rows of the last document put into a table by code (P6), as the window's
            # import would; the table made first when it isn't there
            arg = spec["import_document"]
            if self.last_document is None:
                raise Problem("import_document needs a file_document step before it.")
            tools = Tools(self.world, turn=None)
            if arg["into"] not in self.world.collections.names():
                tools.collection_create(arg["into"], arg.get("title") or arg["into"],
                                        arg["fields"], module=arg.get("module"))
            imported: Any = tools.world_file("import", ref=self.last_document["id"],
                                             into=arg["into"], sheet=arg.get("sheet"))
            record = {"import_document": arg["into"], "ok": "error" not in imported,
                      "detail": str(imported)[:300]}
        elif "wake_agents" in spec:
            # what the scheduler's tick does for the person's agents with on: lines (P7)
            woke = []
            for agent in self.world.agents.working():
                if agent["kind"] == "custom" and agent["events"]:
                    if look_and_run(self.world, "custom", runner=self.runner,
                                    agent_id=agent["id"]) is not None:
                        woke.append(agent["name"])
            record = {"wake_agents": woke, "ok": True,
                      "detail": f"woke {', '.join(woke)}" if woke else "nothing woke"}
        elif "tools_change" in spec:
            # Alpha's hands change between two turns, as an update would (A12): open sessions
            # start afresh, carrying what was going on
            engine.tools_fingerprint = f"{engine.tools_fingerprint}+{new_id('t')}"
            record = {"tools_change": True, "ok": True}
        elif "connect_contacts" in spec:
            from alpha.connectors.contacts import MacContacts

            conn = MacContacts(self.world).connect()
            record = {"connect_contacts": conn["status"], "ok": conn["status"] == "connected"}
        elif "decide_staged" in spec:
            # the person's yes or no on the newest batch an agent brought in (P3 item 5)
            arg = spec["decide_staged"]
            held = self.world.store.one("SELECT id FROM staged WHERE state = 'held' ORDER BY"
                                        " created_at DESC LIMIT 1")
            if held is None:
                record = {"decide_staged": arg, "ok": False, "detail": "Nothing was held."}
            else:
                staging.decide(self.world, held["id"], accept=bool(arg.get("accept")),
                               reason=arg.get("reason"))
                record = {"decide_staged": arg, "ok": True,
                          "detail": f"{held['id']}: {self.world.staged.get(held['id'])['state']}"}
        elif "answer_ask" in spec:
            # the person answers the newest open question matching on its card, through the
            # app's own route (Q92: one a run asked goes back to the agent that asked it)
            arg = spec["answer_ask"]
            pattern = re.compile(str(arg.get("matches") or "."), re.I)
            asks = [a for a in self.world.journal.open_asks() if pattern.search(a["text"])]
            if not asks:
                record = {"answer_ask": arg["text"], "ok": False,
                          "detail": "No open question to answer."}
            else:
                app = TestClient(create_app(self.world, live=False, runner=self.runner))
                res = app.post(f"/api/asks/{asks[-1]['id']}/answer", json={"text": arg["text"]})
                record = {"answer_ask": asks[-1]["text"][:160], "ok": res.status_code == 200,
                          "detail": res.text[:300]}
        elif "break_reader" in spec:
            # A reader broken on purpose on the copy (its script returns nothing), so the
            # journey can watch Doc repair it with the real model (Q49, 16 Oct).
            name = str(spec["break_reader"])
            reader = self.world.readers.get(name)
            self.world.readers.save(name, site=reader["site"], url=reader["url"],
                                    script="return []", description=reader["description"],
                                    to_end=bool(reader.get("to_end")), count=0,
                                    whole=bool(reader.get("whole", True)), by="journey")
            record = {"break_reader": name, "ok": True}
        elif "seed" in spec:
            # Something that was true before the journey began: a past exchange, a fact.
            seed = spec["seed"]
            if seed.get("kind") == "turn":
                # `day: tuesday` is the last such day before today; else `ago: 2d`
                days = (_days_since(str(seed["day"])) if seed.get("day")
                        else int(str(seed.get("ago", "1d")).rstrip("d")))
                at = (dt.now(UTC) - timedelta(days=days)).replace(microsecond=0).isoformat()
                said = self.world.journal.append("said", seed["said"], actor="person")
                replied = self.world.journal.append("replied", seed["replied"],
                                                    data={"turn": said})
                with self.world.store.tx() as db:
                    db.execute("UPDATE journal SET at = ? WHERE id IN (?, ?)", (at, said, replied))
                record = {"seed": f"{days} day(s) ago: {seed['said'][:60]}", "ok": True}
            elif seed.get("kind") == "agent":
                # an agent of the person's, with a brief, there before the journey (P2)
                from_tools = Tools(self.world, turn=None)
                made = from_tools.agent_create(seed["title"], seed.get("schedule", "weekly mon"
                                                                   " 09:00"),
                                               brief=seed["brief"],
                                               name=seed.get("name") or seed["title"],
                                               stage=[self._table(t) for t in
                                                      seed.get("stage") or []] or None,
                                               on=[{k: self._table(v) if k in (
                                                   "rows_added", "rows_changed") else v
                                                   for k, v in line.items()}
                                                   for line in seed.get("wakes_on") or []] or None)
                if "error" in made:
                    raise Problem(made["error"])
                record = {"seed": f"agent {made['name']}", "ok": True}
            elif seed.get("kind") == "table":
                # a table there before the journey, with what its rows are (P8) and its rows
                tools = Tools(self.world, turn=None)
                shaped = tools.world_shape(
                    seed["name"], seed.get("title") or seed["name"], seed["fields"],
                    module=seed.get("module"), title_field=seed.get("title_field"),
                    rows_are=seed.get("rows_are"), identity_field=seed.get("identity_field"),
                    organisation_field=seed.get("organisation_field"),
                    date_field=seed.get("date_field"),
                    direction_field=seed.get("direction_field"))
                if "error" in shaped:
                    raise Problem(shaped["error"])
                rows = [{k: (_days_from_now(v) if isinstance(v, str) and v.startswith("days:")
                             else v) for k, v in row.items()} for row in seed.get("rows") or []]
                if rows:
                    tools.world_write(seed["name"], rows, source="stated")
                record = {"seed": f"table {seed['name']} with {len(rows)} row(s)", "ok": True}
            elif seed.get("kind") == "meeting":
                # a meeting on the calendar before the journey, with its people (P8)
                attendees = []
                for a in seed.get("with") or []:
                    who = self.world.entities.resolve("person", a.get("name") or a["email"],
                                                      {"email": a["email"]})["entity"]
                    attendees.append({"name": a.get("name"), "email": a["email"],
                                      "entity_id": who["id"]})
                at = _days_from_now(f"days:{int(seed.get('days', -1))}", time=True)
                event = self.world.entities.resolve("event", seed["title"],
                                                    {"uid": new_id("uid")})["entity"]
                with self.world.store.tx() as db:
                    db.execute("INSERT INTO events (id, entity_id, connection, uid, calendar,"
                               " title, starts_at, ends_at, all_day, attendees, updated_at)"
                               " VALUES (?,?,NULL,?,'Work',?,?,?,0,?,?)",
                               (new_id("ev"), event["id"], new_id("uid"), seed["title"], at, at,
                                json.dumps(attendees), at))
                record = {"seed": f"meeting {seed['title']}", "ok": True}
            elif seed.get("kind") == "fact":
                self.world.knowledge.record_fact("person", seed["predicate"], seed["value"],
                                                 source="stated", state="accepted")
                record = {"seed": f"fact {seed['predicate']} = {seed['value']}", "ok": True}
            elif seed.get("kind") == "page":
                # a page the model drew, there before the journey (A7)
                self.world.pages.keep(seed["module"], seed["name"], seed.get("title") or
                                      seed["name"], str(seed["html"]), seed.get("data") or {})
                record = {"seed": f"page {seed['name']} in {seed['module']}", "ok": True}
            elif seed.get("kind") == "skill_file":
                # a skill as a file in its module's skills/ (A9), linked where the SDK loads it
                m = self.world.modules.get(seed["module"])
                folder = self.world.files.dir_of(m["id"]) / "skills" / str(seed["name"])
                folder.mkdir(parents=True, exist_ok=True)
                (folder / "SKILL.md").write_text(
                    f"---\nname: {seed['name']}\nkind: {seed.get('skill_kind', 'act')}\n"
                    f"description: {seed['description']}\nmodule: {m['id']}\n---\n"
                    f"{str(seed['body']).strip()}\n")
                self.world.files.sync_module(self.world.modules.get(m["id"]))
                record = {"seed": f"skill {seed['name']} in {m['name']}", "ok": True}
            elif seed.get("kind") == "network":
                # hosts scripts may reach (A3): a module's list, or the world's
                if seed.get("module"):
                    mid = self.world.modules.get(seed["module"])["id"]
                    grants.set_module_hosts(self.world, mid, list(seed["hosts"]), actor="person")
                else:
                    grants.set_world_hosts(self.world, list(seed["hosts"]))
                record = {"seed": f"network {', '.join(seed['hosts'])}", "ok": True}
            elif seed.get("kind") == "note":
                # a note the person wrote (their standing instructions are ALPHA.md's), whole
                self.world.knowledge.write_note(seed.get("scope", "person"), seed["title"],
                                                str(seed["body"]).strip(), source="stated")
                record = {"seed": f"note {seed['title']} ({len(str(seed['body']))} characters)",
                          "ok": True}
            else:
                raise Problem(f"Unknown seed: {seed}")
            # Seeds are not the journey's own work: move the mark past them.
            self.mark = mark(self.world)
        elif "remove_module" in spec:
            # The world as it would be without something the person already has, for a
            # journey that asks for it afresh; not the journey's own work, so the mark moves.
            # A world that never had it is already that world.
            try:
                counts = purge.remove_module(self.world, spec["remove_module"])
                removed: dict[str, Any] = {k: v for k, v in counts.items() if k != "module"}
            except Problem:
                removed = {"already absent": True}
            record = {"remove_module": spec["remove_module"], "ok": True, "removed": removed}
            self.mark = mark(self.world)
        elif "due_now" in spec:
            # the suite cannot wait until Friday: every open reminder's moment is now
            moved = 0
            for c in self.world.tasks.all():
                if c.get("state") not in ("Done", "Dropped") and c.get("remind_at"):
                    self.world.tasks.snooze(c["id"], dt.now(UTC))
                    moved += 1
            record = {"due_now": moved, "ok": moved > 0,
                      "detail": f"{moved} reminder(s) brought to now"}
        elif "connect_folder" in spec:
            # a folder the person connected, with the files the journey names in it
            from alpha.connectors.files import Files

            folder = Path(self.world.path).parent / "granted" / str(spec["connect_folder"]
                                                                   .get("name") or "notes")
            folder.mkdir(parents=True, exist_ok=True)
            for name, text in (spec["connect_folder"].get("files") or {}).items():
                (folder / str(name)).write_text(str(text))
            out = Files(self.world).watch(str(folder))
            record = {"connect_folder": str(folder), "ok": True,
                      "detail": f"{out.get('files', out.get('synced', '?'))} file(s) read"}
        elif "kit_export" in spec:
            # the newest module (or the one named) as a kit, into the scratch home's kits/
            ref = spec["kit_export"].get("module") or self._newest_module()["id"]
            out = self.world.kits.export(ref)
            self.last_kit = out["kit"]
            record = {"kit_export": out["name"], "ok": True,
                      "detail": f"v{out['version']}, {out['files']} files"
                                + (f"; crosses: {out['crosses']}" if out["crosses"] else "")}
        elif "kit_install" in spec:
            path = spec["kit_install"].get("from") or getattr(self, "last_kit", None)
            if not path:
                raise Problem("kit_install needs 'from', or a kit_export step before it.")
            out = self.world.kits.install(Path(path))
            record = {"kit_install": out["name"], "ok": True, "detail": out["words"]}
        elif "module_remove" in spec:
            from alpha.world.purge import remove_module

            ref = spec["module_remove"].get("module") or self._newest_module()["id"]
            gone = remove_module(self.world, ref)
            record = {"module_remove": gone["module"], "ok": True,
                      "detail": f"{gone.get('tables', 0)} tables, {gone.get('agents', 0)} agents"}
        elif "check" in spec:
            # A check between steps (where a sentence went, what exists so far), recorded
            # as a step so the report reads in order.
            out = self.check(spec["check"])
            record = {"check": out["check"], "ok": out["ok"], "detail": out["why"]}
        elif "approve_action" in spec:
            pending = [a for a in self.world.actions.all(("proposed",))
                       if a["created_at"] >= self.mark.at]
            if not pending:
                raise Problem("No action was proposed in this journey to approve.")
            action = sorted(pending, key=lambda a: a["created_at"])[-1]
            done = acting.approve(self.world, action["id"], "Approved by the journey suite",
                                  runner=self.runner, repair=False)
            self.last_action = self.world.actions.get(action["id"])
            record = {"approve_action": action["title"], "ok": bool(done.get("ok")),
                      "detail": done.get("text") or done.get("why")}
        elif "agree_workflow" in spec:
            # The person's yes on the page, as the app's route does it; Alpha then does the
            # work in the conversation, as one turn (harness 2a).
            new = [w for w in self.world.workflows.all(("drafted",))
                   if w["created_at"] >= self.mark.at]
            if not new:
                raise Problem("No workflow was drafted in this journey to agree.")
            agreed = self.world.workflows.agree(new[-1]["id"], "Agreed by the journey suite")
            doing = turn.ask(
                self.world, f"The person agreed the page \"{agreed['title']}\". Do the work"
                            " now, here, as one piece: make what the page says, run it once,"
                            " and report what it holds.",
                thread=agreed.get("conversation"), actor="alpha", runner=self.runner,
                conversation=True, journal_as=f"Agreed: {agreed['title']}. Doing it now.")
            final = self.world.workflows.get(agreed["id"])
            record = {"agree_workflow": agreed["title"], "ok": final["state"] == "built",
                      "state": final["state"], "reply": doing.reply}
        elif "press" in spec:
            # a press on a page's button, as the window passes it (A7): the words as the
            # person's, to the desk with the page's module (section C), journaled as from the
            # page
            a = spec["press"]
            m = self.world.modules.get(a["module"])
            title = next((p["title"] for p in self.world.pages.all(m["id"])
                          if p["name"] == a["page"]), a["page"])
            home = self.world.modules.home_thread()["id"]
            outcome = turn.ask(self.world, a["words"], module=m["id"], thread=home,
                               runner=self._desk_runner(), conversation=True,
                               journal_as=f"From the {title} page: {a['words']}")
            self.last_turn, self.last_reply = outcome.said, outcome.reply
            self.last_thread = home
            record = {"press": a["words"], "ok": outcome.ok, "reply": outcome.reply[:600]}
        elif "add_row" in spec:
            # A thing added as the person would, then the parts that run when it arrives.
            # `table` names it; `field` finds the newest module's table that has that field
            # (the page names things in its own words).
            a = spec["add_row"]
            table = self._table(a["table"]) if a.get("table") else self._table_with(a["field"])
            row = self.world.collections.add(table, a["values"], {"by": "person"})
            self.last_row = (table, str(row["id"]))
            record = {"add_row": table, "ok": True}
        elif "decline_workflow" in spec:
            new = [w for w in self.world.workflows.all(("drafted", "agreed"))
                   if w["created_at"] >= self.mark.at]
            if not new:
                raise Problem("No workflow was drafted in this journey to decline.")
            out = self.world.workflows.decline(new[-1]["id"])
            record = {"decline_workflow": out["title"], "ok": out["state"] == "declined"}
        else:
            raise Problem(f"Unknown step: {spec}")
        record["seconds"] = round(time.monotonic() - began, 1)
        return record

    def _answer_permissions(self, stop: threading.Event, desk_only: bool = False) -> None:
        answered: set[str] = set()
        home = self.world.modules.home_thread()["id"]
        while not stop.wait(0.2):
            for a in self.world.journal.recent(5, kinds=["asked"]):
                if desk_only and a["thread"] != home:
                    continue
                if a["data"].get("permission") == "network" and a["id"] not in answered:
                    # a host's card holds nothing (Q109): the yes lists it, and the desk is
                    # told after the turn, as the app does
                    answered.add(a["id"])
                    self.world.journal.append("answered", "Yes (the journey suite)",
                                              actor="person", data={"ask": a["id"]},
                                              thread=a["thread"], module=a["module"])
                    host = engine.host_answered(self.world, a, "Yes, always")
                    if host and not self.world.agents.of_thread(str(a["thread"] or "")):
                        self.allowed_hosts.append((host, a["module"]))
                    continue
                if a["data"].get("permission") and a["id"] not in answered:
                    answered.add(a["id"])
                    if engine.answer_permission(a["id"], "Yes"):
                        self.world.journal.append("answered", "Yes (the journey suite)",
                                                  actor="person", data={"ask": a["id"]},
                                                  thread=a["thread"], module=a["module"])

    def _file_document(self, arg: dict[str, Any]) -> dict[str, Any]:
        from alpha.connectors.files import Files

        folder = Path(self.world.path).parent / "dropped"
        folder.mkdir(parents=True, exist_ok=True)
        name = str(arg["name"])
        path = folder / name
        if arg.get("again"):
            first = folder / str(arg["again"])
            path.write_bytes(first.read_bytes())
        elif "text" in arg:
            path.write_text(str(arg["text"]))
        elif "sheets" in arg:
            from openpyxl import Workbook

            wb = Workbook()
            for i, (title, rows) in enumerate(dict(arg["sheets"]).items()):
                ws = wb.active if i == 0 else wb.create_sheet()
                assert ws is not None
                ws.title = str(title)
                for row in rows:
                    ws.append(list(row))
            wb.save(path)
        else:
            path.write_bytes(text_pdf([str(p) for p in arg["pages"]]))
        module = None
        if arg.get("module"):
            try:
                module = self.world.modules.get(arg["module"])["id"]
            except Problem:
                module = Tools(self.world, turn=None).module_create(
                    arg["module"], arg.get("goal") or "")["id"]
        return Files(self.world).take(path, module=module, origin="the person",
                                      by="person")

    def _table(self, name: str) -> str:
        names = self.world.collections.names()
        if name in names:
            return name
        hits = [n for n in names if n.endswith("_" + name)]
        if len(hits) != 1:
            raise Problem(f"{len(hits)} tables match '{name}'.")
        return hits[0]

    def _table_with(self, field: str) -> str:
        new = [m for m in self.world.modules.all() if m["id"] not in self.mark.modules]
        for m in reversed(new):
            for t in self.world.collections.overview(m["id"]):
                desc = self.world.collections.describe(t["name"])
                if any(f["name"] == field for f in desc["fields"]):
                    return str(t["name"])
        raise Problem(f"No table of a module made in this journey has a field '{field}'.")

    def _newest_module(self) -> dict[str, Any]:
        modules = sorted(self.world.modules.all(), key=lambda m: m["created_at"])
        if not modules:
            raise Problem("No module yet.")
        return modules[-1]

    def check_kit_installed(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """A kit installed in this world: its module there with its tables (empty) and its
        agents paused."""
        kits = self.world.kits.installed()
        if not kits:
            return False, "No kit is installed."
        k = kits[-1]
        m = self.world.modules.get(str(k["module"]))
        tables = self.world.collections.overview(m["id"])
        agents = self.world.agents.all(m["id"])
        rows = sum(t["records"] for t in tables)
        paused = all(not a["on"] for a in agents)
        skills = [sk for sk in self.world.skills.index() if sk.get("module") == m["id"]]
        words = (f"{k['kit']} v{k['version']}: module {m['name']} with {len(tables)} table(s),"
                 f" {rows} row(s), {len(skills)} skill(s), {len(agents)} agent(s)"
                 f"{' (paused)' if paused else ''}")
        ok = (bool(tables) and (not arg.get("empty", True) or rows == 0) and paused
              and len(skills) >= int(arg.get("skills", 0)))
        return ok, words

    def _agent(self, title: str) -> dict[str, Any]:
        hits = [a for a in self.world.agents.all()
                if title.lower() in a["name"].lower() or title.lower() in a["description"].lower()]
        if len(hits) != 1:
            raise Problem(f"{len(hits)} agents match '{title}'.")
        return hits[0]

    # checks

    def check(self, spec: dict[str, Any]) -> dict[str, Any]:
        (kind, arg), = spec.items()
        began = time.monotonic()
        method = getattr(self, f"check_{kind}", None)
        if method is None:
            raise Problem(f"Unknown check: {kind}")
        try:
            ok, why = method(arg if arg is not None else {})
        except Problem as e:
            ok, why = False, str(e)
        return {"check": kind, "arg": arg, "ok": ok, "why": why,
                "seconds": round(time.monotonic() - began, 1)}

    def _new_rows(self, collection: str) -> list[dict[str, Any]]:
        return [r for r in self.world.collections.query(collection, limit=None)
                if r["created_at"] >= self.mark.at]

    def check_independent(self, arg: dict[str, Any]) -> tuple[bool, str]:
        if not self.last_turn:
            return False, "No turn to check."
        result = columbo.check_turn(self.world, self.last_turn, runner=self.runner, repair=False)
        if not result["checked"]:
            return False, result["why"]
        return bool(result["agree"]), columbo.words(result)

    def check_row(self, arg: dict[str, Any]) -> tuple[bool, str]:
        rows = self._new_rows(arg["collection"])
        if not rows:
            return False, f"No row was added to {arg['collection']}."
        bad = [s for s in arg.get("source_not", [])
               if any((r.get("_provenance") or {}).get("source") == s for r in rows)]
        if bad:
            return False, f"A new row's source is {', '.join(bad)}."
        sources = sorted({str((r.get("_provenance") or {}).get("source")) for r in rows})
        return True, f"{len(rows)} row(s) added; source: {', '.join(sources)}."

    def check_near(self, arg: dict[str, Any]) -> tuple[bool, str]:
        rows = self._new_rows(arg["collection"])
        values: list[float] = [float(r[arg["field"]]) for r in rows
                               if isinstance(r.get(arg["field"]), int | float)]
        if not values:
            return False, f"No new row in {arg['collection']} has a number in {arg['field']}."
        target, within = float(arg["value"]), float(arg.get("within", 0.1))
        hit = [v for v in values if abs(v - target) <= within * target]
        words = (f"{arg['field']} = {', '.join(str(v) for v in values)}"
                 f" (wanted {target} ±{within:.0%})")
        return bool(hit), words

    def _work(self) -> list[dict[str, Any]]:
        """The work beside the desk this journey started (section C)."""
        return [t for t in self.world.modules.items(limit=200)
                if t["created_at"] >= self.mark.at and t.get("session_tools") == desk.WORK]

    def _outcome_reply(self, thread: str) -> dict[str, Any] | None:
        home = self.world.modules.home_thread()["id"]
        return next((e for e in self.world.journal.recent(200, thread=home)
                     if e["kind"] == "replied" and thread in (e["data"].get("work") or [])),
                    None)

    def check_told(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """Agents' runs that came to the conversation (C5): exactly `count` lines in Home marked
        as told, each with its agent's name; any other run stayed on its page."""
        home = self.world.modules.home_thread()["id"]
        lines = [e for e in self.world.journal.recent(200, thread=home, kinds=["replied"])
                 if e["at"] >= self.mark.at and e["data"].get("told")]
        words = "; ".join(f"{e['data'].get('agent_name')}: {e['text'][:80]}" for e in lines)
        return len(lines) == int(arg.get("count", 1)), f"{len(lines)} told: {words or 'none'}."

    def check_answered_in_words(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """Yes cards the desk answered on the person's words (C4): exactly `count`."""
        got = [e for e in self.world.journal.recent(200, kinds=["answered"])
               if e["at"] >= self.mark.at and e["data"].get("in_words")]
        words = "; ".join(f"{e['text']} (\"{e['data']['in_words']}\")" for e in got) or "none"
        return len(got) == int(arg.get("count", 1)), f"{len(got)} answered in words: {words}."

    def check_asked_which(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """After words that don't say which card, Alpha asks which on a card of its own (C4)."""
        said = self.world.journal.read(self.last_turn) if self.last_turn else None
        asked = [a for a in self.world.journal.open_asks()
                 if not a["data"].get("permission") and said and a["at"] >= said["at"]]
        return bool(asked), ("Asked: " + asked[-1]["text"][:120] if asked
                             else "No question asked after the last sentence.")

    def check_yes_cards_open(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """How many yes cards are still open (`count`): none answered on words that didn't say
        which."""
        still = [a for a in self.world.journal.open_asks() if a["data"].get("permission")]
        return len(still) == int(arg.get("count", 0)), f"{len(still)} yes card(s) open."

    def check_work(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """Work beside the desk (section C, C2's proof): how many pieces ran (`count`), each
        ended with its outcome on its thread and the desk's words on it in Home carrying its
        id; `answered_while_running`: a sentence the person said while the first piece ran,
        other than the one that started it, got its reply (in its own turn, or in the turn
        that also told the work's outcome, when the work finished while it ran: Q113)."""
        work = self._work()
        words = "; ".join(f"{t['title']} ({t['state']})" for t in work) or "none"
        if "count" in arg and len(work) < int(arg["count"]):
            return False, f"{len(work)} piece(s) of work, wanted {arg['count']}: {words}."
        for t in work:
            if not self._outcome_reply(t["id"]):
                return False, f"No reply in Home carries the work {t['title']}: {words}."
        if arg.get("answered_while_running") and work:
            first = min(work, key=lambda t: t["created_at"])
            ended = self._outcome_reply(first["id"])
            home = self.world.modules.home_thread()["id"]
            lines = self.world.journal.recent(200, thread=home)
            began = next((e["data"].get("turn") for e in lines if e["kind"] == "did"
                          and e["data"].get("work") == first["id"]), None)
            meanwhile = [e["id"] for e in lines if e["kind"] == "said"
                         and e["actor"] == "person" and e["id"] != began
                         and first["created_at"] < e["at"] <= (ended or {}).get("at", "~")]
            answered = [e for e in lines if e["kind"] == "replied"
                        and e["data"].get("turn") in meanwhile]
            if not answered:
                return False, f"Nothing said while {first['title']} ran was answered."
            words += f"; answered meanwhile: {answered[0]['text'][:80]}"
        return True, f"{len(work)} piece(s): {words}."

    def check_opened(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """Nothing opened as an item (section C's acceptance): every thread the journey made
        is work beside the desk or an agent's own; `count` is how many others there may be."""
        made = [t for t in self.world.modules.items(limit=200)
                if t["created_at"] >= self.mark.at and t.get("session_tools") != desk.WORK
                and not self.world.agents.of_thread(t["id"])]
        words = "; ".join(f"{t['title']} ({t['state']})" for t in made) or "none"
        return len(made) == int(arg.get("count", 0)), f"{len(made)} opened: {words}."

    def check_made_since(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """What the journey made: at least `tables`, `readers`, `agents`, `modules` (counts),
        whatever Alpha chose to call them (the first-day sentences, harness design §17)."""
        tables = [c for c in self._tables_since()]
        readers = [k for k in self.world.skills.all("read")
                   if k.get("created_at", "") >= self.mark.at]
        agents = [a for a in self.world.agents.all() if a.get("created_at", "") >= self.mark.at]
        modules = [m for m in self.world.modules.all() if m.get("created_at", "") >= self.mark.at]
        have = {"tables": len(tables), "readers": len(readers), "agents": len(agents),
                "modules": len(modules)}
        short = [f"{k}: {have[k]} of {v}" for k, v in arg.items() if have.get(k, 0) < int(v)]
        words = ", ".join(f"{have[k]} {k}" for k in ("tables", "readers", "agents", "modules"))
        return (not short, f"Made {words}" + (f"; short on {', '.join(short)}" if short else ""))

    def check_rows_since(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """The rows in the tables the journey made add up to at least `at_least`."""
        counts = {c["name"]: int((self.world.store.one(
            "SELECT COUNT(*) AS n FROM records WHERE collection = ?",
            (c["name"],)) or {"n": 0})["n"])
            for c in self._tables_since()}
        total = sum(counts.values())
        words = ", ".join(f"{n} ({c})" for n, c in counts.items()) or "no table"
        return total >= int(arg["at_least"]), f"{total} rows in {words}"

    def _tables_since(self) -> list[dict[str, Any]]:
        """The person's tables made in this journey (a module's settings table is Alpha's, and
        the world's own, the tasks, is no table the journey made)."""
        rows = self.world.store.all(
            "SELECT name, created_at FROM collections WHERE created_at >= ?"
            " AND name != 'settings' AND json_extract(schema, '$.world') IS NULL",
            (self.mark.at,))
        return [dict(r) for r in rows]

    def check_no_new_tables(self, arg: dict[str, Any]) -> tuple[bool, str]:
        new = {t["name"] for t in self.world.collections.overview()} - self.mark.tables
        return not new, ("No table was made." if not new else f"Made tables: {sorted(new)}.")

    def check_no_new_skills(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """Alpha used the know-how it had: no reader, procedure or pipeline was written or
        rewritten in this journey (a reuse with the same version is what passes)."""
        current = {s["name"]: s["version"] for s in self.world.skills.all()}
        new = sorted(n for n in current if n not in self.mark.skills)
        changed = sorted(n for n, v in current.items()
                         if n in self.mark.skills and v != self.mark.skills[n])
        if new or changed:
            return False, (f"Wrote new skills: {new}. " if new else "") + \
                (f"Rewrote: {changed}." if changed else "")
        return True, "No skill was written or rewritten."

    def check_no_new_modules(self, arg: dict[str, Any]) -> tuple[bool, str]:
        new = {m["id"] for m in self.world.modules.all()} - self.mark.modules
        return not new, ("No module was made." if not new else f"Made {len(new)} module(s).")

    def check_workflow(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """A workflow page drafted in this journey (§12), in the state asked, with at least
        `steps_min` steps, `keeps_min` things kept, `suggestions_min` suggestions (`sherlock_min`
        of them an agent's, Q45), and, when
        given, `version` (exactly) or `version_min` (a change landed, however many revisions
        the model took to get it right)."""
        new = [w for w in self.world.workflows.all() if w["created_at"] >= self.mark.at]
        if not new:
            return False, "No workflow page was drafted."
        wf = new[-1]
        b = wf["body"]
        by_agents = [sg for sg in b["suggestions"] if sg.get("by")]
        ok = (wf["state"] == arg.get("state", "drafted")
              and len(b["steps"]) >= int(arg.get("steps_min") or 0)
              and len(b["keeps"]) >= int(arg.get("keeps_min") or 0)
              and len(b["suggestions"]) >= int(arg.get("suggestions_min") or 0)
              and len(by_agents) >= int(arg.get("sherlock_min") or 0)
              and (not arg.get("version") or wf["version"] == int(arg["version"]))
              and wf["version"] >= int(arg.get("version_min") or 0))
        steps = "; ".join(f"{s['title']} ({s['who']}{', asks' if s['outward'] else ''})"
                          for s in b["steps"])
        return ok, (f"'{wf['title']}' v{wf['version']} is {wf['state']}: {len(b['steps'])} steps:"
                    f" {steps}. Keeps {', '.join(k['thing'] for k in b['keeps']) or 'nothing'}."
                    f" {len(b['suggestions'])} suggestions ({len(by_agents)} Sherlock's):"
                    f" {'; '.join(_titled(sg) for sg in b['suggestions']) or 'none'}."
                    f" Settings {len(b['settings'])}, questions {len(b['questions'])}.")

    def check_picks(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """Scout's picks for the day (Q48): at least `min` of them on Home, each with evidence
        that resolves and one action."""
        from alpha.runtime.agents import scout

        picks = [p for p in scout.picks_of_day(self.world) if p["at"] >= self.mark.at]
        ok = len(picks) >= int(arg.get("min") or 1)
        words = "; ".join(f"{p['title']} [{p['kind']}{', less sure' if p.get('less_sure') else ''}"
                          f"{', card' if p.get('card') else ''}] → {p['action']}" for p in picks)
        return ok, f"{len(picks)} pick(s): {words or 'none'}."

    def check_ask(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """A question for the person, as a card, from the last turn (Q35): `derived` whether
        the core made it from the reply, `options_min` how many choices at least."""
        if not self.last_turn:
            return False, "No turn."
        asks = [a for a in self.world.journal.open_asks()
                if a["data"].get("turn") == self.last_turn]
        if not asks:
            return False, "The last turn made no card."
        a = asks[-1]
        words = (f"\"{a['text'][:80]}\" with {a['data'].get('options') or 'no'} options"
                 f"{' (made by the core)' if a['data'].get('derived') else ' (the model\'s own)'}")
        if "derived" in arg and bool(a["data"].get("derived")) != bool(arg["derived"]):
            return False, words
        if len(a["data"].get("options") or []) < int(arg.get("options_min", 0)):
            return False, words
        return True, words

    def check_reply(self, arg: dict[str, Any]) -> tuple[bool, str]:
        text = self.last_reply
        if "contains" in arg and arg["contains"].lower() not in text.lower():
            return False, f"The reply doesn't say \"{arg['contains']}\"."
        if "matches" in arg and not re.search(arg["matches"], text, re.I):
            return False, f"The reply doesn't match /{arg['matches']}/."
        return True, "The reply reads as expected."

    def _given(self, turn_id: str, title: str, cap: int = 1_500) -> str:
        """One part of what the model was given for a turn (the layers kept with it)."""
        kept = self.world.journal.context(turn_id)
        return layers.part_of(str(kept["context"]) if kept else "", title)[:cap]

    def check_judge(self, arg: Any) -> tuple[bool, str]:
        rubric = arg if isinstance(arg, str) else arg["rubric"]
        if not self.last_turn:
            return False, "No turn to judge."
        said = self.world.journal.read(self.last_turn)
        system = ("You judge one answer against a rubric. Reply with JSON only:"
                  ' {"pass": true|false, "why": "one sentence"}. What Alpha did in the turn,'
                  " from its own journal, is the evidence of what it read: an answer drawn"
                  " from those steps is not invented.")
        earlier = ""
        if isinstance(arg, dict) and arg.get("history") and len(self.exchanges) > 1:
            shown = self.exchanges[-1 - int(arg["history"]):-1]
            earlier = "Earlier in this conversation:\n" + "\n".join(
                f'Person: "{s}"\nAlpha: {r}' for s, r in shown) + "\n\n"
        # 14 Oct 2026: gmail_network's judge saw only the answer and called four real Gmail
        # reads invented; the turn's steps are what Alpha read and did.
        steps = [s for s in self.world.journal.of_turn(self.last_turn)
                 if s["kind"] in ("saw", "did", "made", "changed", "failed", "noticed")]
        done = ("What Alpha did in this turn, from its journal:\n" + "\n".join(
            f"- {s['kind']}: {' '.join(s['text'].split())[:240]}" for s in steps) + "\n\n"
                if steps else "Alpha's journal records no step in this turn.\n\n")
        # 15 Oct 2026: memory_noticing's judge called a remembered fact (Vikas moving to
        # Bangalore, kept by Watson the turn before) invented, because memory leaves no step
        # in the turn that uses it. What the model was given about whoever the sentence names
        # is the evidence of what it remembered.
        known = self._given(self.last_turn, "WHO THE SENTENCE NAMES")
        # 19 Oct 2026: when_did_we_last_talk's judge called the email subjects and dates
        # invented, though the search's matches given with the sentence held every one
        matched = self._given(self.last_turn, "MATCHES FOR THIS SENTENCE")
        given = "\n".join(part for part in (known, matched) if part)
        remembered = (f"What Alpha already knew when it answered (its memory and the matches"
                      f" for the sentence, given to it with the sentence):\n{given}\n\n"
                      if given else "")
        # 19 Oct 2026: a read in SQL (world_query) leaves no step in the journal, so a true
        # detail read that way looked invented; a journey that set its world up says what is
        # true there, and the answer is judged against that
        # 19 Oct 2026, night (section B): with the person's card no longer given with the
        # sentence, the judge still called a true count "not grounded": an answer that agrees
        # with what is true was read, even when the read left no step
        truth = (f"What is true in the person's world (the journey set it up; today is"
                 f" {dt.now(UTC).date().isoformat()}):\n{arg['facts']}\nAlpha reads its tables"
                 " in SQL, and a read leaves no step in the journal: an answer that agrees with"
                 " what is true here was read, not invented.\n\n"
                 if isinstance(arg, dict) and arg.get("facts") else "")
        # 20 Oct 2026 (C6): what a piece of work beside the desk did is its own journal, its
        # steps and its outcome; the desk holds only the outcome and reads the steps with one
        # query, which leaves no step either. An answer about the work is judged against them.
        worked = ""
        if isinstance(arg, dict) and arg.get("stream"):
            # section C's acceptance: the one stream as the person read it, in its order
            home = self.world.modules.home_thread()["id"]
            worked += "The conversation, oldest first:\n" + "\n".join(
                f"- {e['at'][11:19]} {'Person' if e['kind'] == 'said' else 'Alpha'}:"
                f" {' '.join(e['text'].split())[:600]}"
                for e in self.world.journal.recent(200, thread=home,
                                                   kinds=["said", "replied", "failed"])
                if e["at"] >= self.mark.at) + "\n\n"
        if isinstance(arg, dict) and arg.get("work_journal"):
            home = self.world.modules.home_thread()["id"]
            for t in self._work():
                # its own journal, and the desk's part in it (starting it, its words on it)
                lines = self.world.journal.recent(200, thread=t["id"]) + [
                    e for e in self.world.journal.recent(200, thread=home)
                    if t["id"] in str(e["data"].get("work") or "")]
                worked += (f"The work \"{t['title']}\" from its own journal, and the"
                           " conversation's lines about it, oldest first:\n")
                worked += "\n".join(
                    f"- {e['at'][11:19]} {e['kind']}: {' '.join(e['text'].split())[:400]}"
                    for e in sorted(lines, key=lambda e: e["at"])) + "\n\n"
        prompt = (f'{earlier}The person said: "{said["text"]}"\n\n{done}{remembered}{truth}'
                  f"{worked}The answer:\n{self.last_reply}\n\nThe rubric: {rubric}")
        result = self.runner(running.TurnRequest(
            sentence=prompt, system=system, world_path=self.world.path,
            turn_id=self.last_turn, kind="judge"))
        if not result.ok:
            return False, f"The judge did not answer: {result.error}"
        m = re.search(r"\{.*\}", result.reply, re.S)
        try:
            verdict = json.loads(m.group(0)) if m else {}
        except json.JSONDecodeError:
            verdict = {}
        return bool(verdict.get("pass")), str(verdict.get("why") or result.reply[:300])

    def check_task(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """A task the journey made (P1): its title matching `matches`, living on the kind of
        place `on` names (record, field, module, page; "nowhere" for none)."""
        rows = self._new_rows(TASKS)
        pattern = re.compile(str(arg.get("matches") or "."), re.I)
        hits = [r for r in rows if pattern.search(str(r.get("title") or ""))]
        if not hits:
            return False, (f"No task matching {arg.get('matches')!r}; made:"
                           f" {[r.get('title') for r in rows] or 'none'}.")
        task = hits[-1]
        steps = self.world.tasks.place(task["anchor"]) if task.get("anchor") else []
        kind = steps[-1]["kind"] if steps else "nowhere"
        where = " › ".join(s["words"] for s in steps) or "no place"
        want = arg.get("on")
        return (want is None or kind == want), f"{task['title']}: on {where} ({kind})."

    def check_task_run(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """A task the journey made was served by an agent's run (P2): the run's verdict, the
        task's state after it, and steps that carry the run."""
        pattern = re.compile(str(arg.get("matches") or "."), re.I)
        tasks = [t for t in self._new_rows(TASKS) if pattern.search(str(t.get("title") or ""))]
        if not tasks:
            return False, f"No task matching {arg.get('matches')!r}."
        task = self.world.tasks.get(tasks[-1]["id"])
        runs = self.world.runs.of_task(task["id"])
        if not runs:
            return False, f"{task['title']}: no run served it (state {task['state']})."
        run = runs[0]
        steps = self.world.runs.steps(run["id"])
        ok = (run["verdict"] == arg.get("verdict", "succeeded")
              and task["state"] == arg.get("state", "Done") and bool(steps))
        said = (run.get("line") or run.get("why") or "")[:160]
        return ok, (f"{task['title']}: run {run['verdict']} ({run['trigger']}), {len(steps)}"
                    f" steps, task {task['state']}: {said}")

    def check_card(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """The newest card the journey made (P3): its `kind`, its `state`, and with `by_agent`
        drafted by one of the person's agents with the task its run served."""
        rows = [c for c in self.world.cards.decided(200) if c["created_at"] >= self.mark.at
                and (not arg.get("kind") or c["kind"] == arg["kind"])]
        if not rows:
            return False, f"No {arg.get('kind') or ''} card was made.".replace("  ", " ")
        c = rows[0]
        by = c["drafter"] not in ("alpha", "person")
        ok = ((not arg.get("state") or c["state"] == arg["state"])
              and (not arg.get("by_agent") or (by and bool(c.get("task")))))
        return ok, (f"{c['kind']} card \"{c['text'][:100]}\", {c['state']}, drafted by"
                    f" {c['drafter']}" + (f", for task {c['task']}" if c.get("task") else ""))

    def check_agent_made(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """The agent the journey made in words (Q92): a short name, not the sentence; on call
        (no schedule) when `on_call`; and no task of its own."""
        rows = self.world.store.all("SELECT id FROM agents WHERE kind = 'custom' AND"
                                    " created_at >= ? ORDER BY created_at", (self.mark.at,))
        if not rows:
            return False, "No agent was made."
        agent = self.world.agents.get(rows[-1]["id"])
        words = len(agent["name"].split())
        owned = [t["title"] for t in self.world.tasks.of(agent["id"], finished=True)]
        ok = (words <= int(arg.get("name_words_max", 4)) and not owned
              and (not arg.get("on_call") or agent["schedule"] is None))
        return ok, (f"{agent['name']!r} ({words} words), works {agent['when']}; tasks it owns:"
                    f" {owned or 'none'}.")

    def check_answer_went_back(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """The person's answer to a question a run asked went back to the agent that asked
        (Q92): a run of it, started by the answer; nothing about the person kept from it."""
        answers = [e for e in self.world.journal.recent(50, kinds=["answered"])
                   if e["at"] >= self.mark.at]
        if not answers:
            return False, "Nothing was answered."
        last = answers[-1]
        to = last["data"].get("to")
        if not to:
            return False, f"The answer {last['text'][:80]!r} went to Alpha's conversation."
        runs = [r for r in self.world.runs.of(str(to)) if r["trigger"] == "answer"]
        if not runs:
            return False, "No run of the agent took the answer."
        run = runs[0]
        kept = [e["text"] for e in self.world.journal.recent(50, kinds=["noticed"])
                if e["at"] >= last["at"] and e["data"].get("subject") == "person"]
        ok = not kept and ("verdict" not in arg or run["verdict"] == arg["verdict"])
        return ok, (f"{self.world.agents.get(str(to))['name']}'s run {run['verdict']}:"
                    f" {(run.get('line') or run.get('why') or '')[:200]}"
                    + (f"; kept about the person: {kept}" if kept else ""))

    def check_value_source(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """The values of a field say where they came from (P5): every row's source is of the
        kind asked (a document's cell), and with `row`/`cell`, that row's is that cell."""
        table = self._table(arg["collection"])
        rows = [r for r in self.world.collections.query(table, limit=None)
                if r.get(arg["field"]) not in (None, "")]
        if not rows:
            return False, f"No row of {table} has {arg['field']}."
        sources = [((r.get("_provenance") or {}).get("fields") or {}).get(arg["field"]) or {}
                   for r in rows]
        kind = arg.get("kind", "document")
        wrong = [r for r, src in zip(rows, sources, strict=True) if src.get("kind") != kind]
        words = "; ".join(f"{r.get(arg['field'])} from {src.get('source') or 'no source'}"
                          for r, src in list(zip(rows, sources, strict=True))[:4])
        if wrong:
            return False, f"{len(wrong)} of {len(rows)} without a {kind} source: {words}."
        if arg.get("cell"):
            hits = [src for r, src in zip(rows, sources, strict=True)
                    if str(r.get(arg["match_field"])) == str(arg["row"])]
            if not hits or hits[0].get("cell") != arg["cell"]:
                return False, f"{arg['row']}'s {arg['field']} isn't from {arg['cell']}: {words}."
        return True, f"{len(rows)} value(s), each from its place: {words}."

    def check_kept_once(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """One file, one document (P6): the same bytes taken twice are one document, and the
        second taking said it was already kept."""
        rows = self.world.store.all("SELECT title FROM documents WHERE removed_at IS NULL"
                                    " AND title LIKE ?", (f"%{arg['title_contains']}%",))
        again = bool(self.last_document and self.last_document.get("already"))
        return len(rows) == 1 and again, (f"{len(rows)} document(s) titled like"
                                          f" {arg['title_contains']!r}; the second taking"
                                          f" {'said already kept' if again else 'made a copy'}.")

    def check_woke(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """An agent woken by events (P7): that many runs started by them, and the first line
        of the newest names what woke it."""
        agent = self._agent(arg["agent"])
        runs = [r for r in self.world.runs.of(agent["id"]) if r["trigger"] == "event"]
        lines = [e["text"] for e in self.world.journal.recent(50)
                 if e["thread"] == agent["thread"] and "woken by" in e["text"]]
        said = lines[-1] if lines else ""
        ok = len(runs) == int(arg.get("runs", 1)) and (
            not arg.get("contains") or str(arg["contains"]).lower() in said.lower())
        return ok, (f"{len(runs)} run(s) woken by events" + (f"; {said[:240]}" if said else "")
                    + (f"; last: {(runs[0].get('line') or runs[0].get('why') or '')[:160]}"
                       if runs else ""))

    def _person(self, arg: dict[str, Any]) -> dict[str, Any]:
        keys = {k: arg[k] for k in ("email", "phone", "url") if arg.get(k)}
        found = self.world.entities.find(keys=keys)
        if not found:
            raise Problem(f"No one is known by {keys}.")
        return found[0]

    def check_one_person(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """Everything the journey names (`same`: other keys) is one person (P8), with
        `sources` among where they came from and rows of `in_table`."""
        person = self._person(arg)
        others = [self._person(dict(other)) for other in arg.get("same") or []]
        apart = [o["name"] for o in others if o["id"] != person["id"]]
        sources = {r["source"].split(":")[0] for r in people_world.sources_of(
            self.world, person["id"])}
        rows = {r["collection"] for r in self.world.collections.linked_to(person["id"])}
        wanted = set(arg.get("sources") or [])
        table = arg.get("in_table")
        ok = not apart and wanted <= sources and (
            not table or any(t == table or t.endswith(f"_{table}") for t in rows))
        reach = ", ".join(v for k, vs in person["keys"].items() if k in ("email", "phone")
                          for v in vs)
        return ok, (f"{person['name']} ({reach}); sources {sorted(sources) or 'none'};"
                    f" rows in {sorted(rows) or 'no table'}"
                    + (f"; apart: {apart}" if apart else "") + ".")

    def check_pairs(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """Possible same people wait on a card and none was merged by code (P8)."""
        rows = self.world.store.all("SELECT * FROM maybe_same WHERE state = 'proposed'")
        names = []
        for r in rows:
            try:
                names.append((self.world.entities.get(r["a"])["name"],
                              self.world.entities.get(r["b"])["name"], r["why"]))
            except Problem:
                continue
        named = str(arg.get("named") or "")
        hits = [n for n in names if not named or (n[0] == named and n[1] == named)]
        merged = self.world.store.all("SELECT 1 FROM maybe_same WHERE state = 'merged'")
        ok = len(hits) >= int(arg.get("at_least", 1)) and not merged
        return ok, (f"{len(rows)} pair(s) waiting: " + "; ".join(
            f"{a} / {b} ({why})" for a, b, why in names[:5]) + (f"; {len(merged)} merged"
                                                                  if merged else "") + ".")

    def check_organisation(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """The organisation of a domain, with at least that many people (P8)."""
        found = self.world.entities.find(keys={"domain": arg["domain"]})
        if not found:
            return False, f"No organisation for {arg['domain']}."
        org = found[0]
        members = self.world.entities.members(org["id"])
        ok = len(members) >= int(arg.get("people_at_least", 1))
        return ok, (f"{org['name']} ({arg['domain']}): "
                    + ", ".join(m["name"] for m in members) + ".")

    def check_people(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """How many people Alpha knows now (P9)."""
        row = self.world.store.one("SELECT COUNT(*) AS n FROM entities WHERE kind = 'person'"
                                   " AND merged_into IS NULL")
        n = int(row["n"]) if row else 0
        return n >= int(arg.get("at_least", 1)), f"{n} people."

    def check_count(self, arg: dict[str, Any]) -> tuple[bool, str]:
        rows = self.world.collections.query(self._table(arg["collection"]), limit=None)
        n = len(rows)
        if "reader" in arg:
            rows = [r for r in rows if r.get("_gone_at") is None]
            n = len(rows)
            reader = self.world.readers.get(arg["reader"])
            fraction = float(arg.get("at_least_fraction_of_last_ok", 0.95))
            floor = (reader.get("last_ok_count") or 0) * fraction
            return n >= floor, f"{n} rows kept; the reader last read {reader.get('last_ok_count')}."
        want = int(arg.get("at_least", 1))
        return n >= want, f"{n} rows (wanted at least {want})."

    def check_reader_health(self, arg: Any) -> tuple[bool, str]:
        reader = self.world.readers.get(arg if isinstance(arg, str) else arg["name"])
        return reader["health"] == "ok", (
            f"{reader['name']}: {reader['health']}, {reader.get('last_count')} rows"
            + (f"; {reader['last_problem']}" if reader.get("last_problem") else "") + ".")

    def check_run_verdict(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """The last run's verdict, judged by code, is one of the allowed."""
        auto = self.last_agent or self._agent(arg["title"])
        last = self.world.runs.last(auto["id"])
        if last is None:
            return False, "No run ended."
        allowed = arg.get("one_of") or [arg.get("verdict", "succeeded")]
        words = f"{last['verdict']}: {last['why'] or last['line']}" + (
            f" (model {last['model_ms'] // 1000} s, {last['repairs']} repairs)")
        return last["verdict"] in allowed, words

    def check_agent_is_pipeline(self, arg: Any) -> tuple[bool, str]:
        """The agent runs as steps now (no model), not by its brief."""
        if self.last_agent:
            auto = self.last_agent
        elif isinstance(arg, str) or (isinstance(arg, dict) and arg.get("title")):
            auto = self._agent(arg if isinstance(arg, str) else arg["title"])
        else:
            # the newest agent the journey made, whatever Alpha named it
            made = [a for a in self.world.agents.all() if a.get("created_at", "") >= self.mark.at]
            if not made:
                return False, "No agent was made in this journey."
            auto = made[-1]
        auto = self.world.agents.get(auto["id"])
        steps = auto.get("steps") or []
        reads = [s["read"] for s in steps if "read" in s]
        return bool(steps), (f"Steps: {', '.join(reads)} then tell." if steps
                             else "Still a brief the model follows.")

    def check_agent(self, arg: dict[str, Any]) -> tuple[bool, str]:
        auto = self.last_agent or self._agent(arg["title"])
        result, problem = auto.get("last_result") or "", auto.get("last_error") or ""
        if "result_matches" in arg and not re.search(arg["result_matches"], result):
            return False, f"Result \"{result[:120]}\" doesn't match /{arg['result_matches']}/."
        if problem and not (arg.get("problem_allowed")
                            and re.search(arg["problem_allowed"], problem, re.I)):
            return False, f"Problem: {problem}"
        return True, f"{result[:160]}" + (f" Problem allowed: {problem}" if problem else "")

    def check_action(self, arg: dict[str, Any]) -> tuple[bool, str]:
        new = [a for a in self.world.actions.all() if a["created_at"] >= self.mark.at]
        if not new:
            return False, "No action was proposed."
        action = self.last_action or sorted(new, key=lambda a: a["created_at"])[-1]
        action = self.world.actions.get(action["id"])
        words = (f"\"{action['title']}\" ({action['effect']} on {action['site']}):"
                 f" {action['state']}"
                 + (f"; {action['error']}" if action.get("error") else "")
                 + (f"; preview {Path(action['preview']).name}" if action.get("preview") else ""))
        if "state" in arg and action["state"] != arg["state"]:
            return False, words
        if "effect" in arg and action["effect"] != arg["effect"]:
            return False, words
        if arg.get("preview") and not (action.get("preview") and Path(action["preview"]).exists()):
            return False, words + "; no preview screenshot"
        if "procedure" in arg and action["procedure"] != arg["procedure"]:
            return False, words + f"; through {action['procedure']}, not {arg['procedure']}"
        return True, words

    def check_document(self, arg: dict[str, Any]) -> tuple[bool, str]:
        module = self.world.modules.get(arg["module"])["id"] if arg.get("module") else None
        rows = [r for r in self.world.store.all(
            "SELECT * FROM documents WHERE removed_at IS NULL AND indexed_at >= ?"
            " ORDER BY indexed_at", (self.mark.at,))
            if (module is None or r["module"] == module)
            and (not arg.get("kind") or r["kind"] == arg["kind"])]
        if not rows:
            return False, "No such document was kept."
        return True, ", ".join(f"{r['title']} ({r['size']} bytes, {len(r['text'].split())} words)"
                               for r in rows) + "."

    def check_page(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """A page the model drew (slice 7): in the module, with a header that reads, its data
        running and (when asked) that many rows in a named set, the module's front when
        asked, and a drawing of the kind asked (an <svg> or a <canvas>) in its text."""
        pages = [p for p in self.world.pages.all(arg["module"]) if not p["problem"]]
        broken = [p for p in self.world.pages.all(arg["module"]) if p["problem"]]
        if not pages:
            return False, ("No page draws" + (f": {broken[0]['problem']}" if broken else "") + ".")
        fronts = [p for p in pages if p["front"]]
        page = fronts[0] if arg.get("front") and fronts else pages[0]
        if arg.get("name"):
            named = [p for p in pages if p["name"] == arg["name"]]
            if not named:
                return False, (f"No page named {arg['name']}; there are "
                               + ", ".join(p["name"] for p in pages) + ".")
            page = named[0]
        if arg.get("front") and not fronts:
            return False, f"{len(pages)} page(s), none the module's front."
        if arg.get("count") and len(pages) + len(broken) != int(arg["count"]):
            return False, (f"{len(pages) + len(broken)} pages, not {arg['count']}: "
                           + ", ".join(p["name"] for p in pages + broken) + ".")
        try:
            rows = self.world.pages.data(arg["module"], page["name"])
        except Problem as e:
            return False, f"The page {page['title']}: {e}"
        counts = {k: len(v) for k, v in rows.items()}
        html = self.world.pages.render(arg["module"], page["name"]).lower()
        kit = "alpha.ui.page(" in html
        draws = kit or "<svg" in html or "<canvas" in html
        how = ("draws with " + ("the kit" if kit else "svg" if "<svg" in html else "canvas")) \
            if draws else "no svg, canvas or kit block in it"
        words = (f"{page['title']} ({page['name']}.html){', the front' if page['front'] else ''}:"
                 f" {', '.join(f'{k} {n} rows' for k, n in counts.items()) or 'no data'}; {how}")
        enough = all(n >= int(arg["rows"]) for n in counts.values()) if arg.get("rows") else True
        if arg.get("kit") and not kit:
            words += "; NOT drawn with the kit"
        return bool(counts) and enough and draws and (kit or not arg.get("kit")), words + "."

    def check_page_look(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """Alpha's own look at the module's front page (slice 7b): rendered in Chromium and
        judged by the small model with eyes against the six rules."""
        page = self.world.pages.front(arg["module"])
        if page is None:
            return False, "No front page to look at."
        seen = look.page(self.world, arg["module"], page["name"], runner=self.runner)
        return seen.clean, (f"{page['title']}: " + ("clean" if seen.clean else seen.words)
                            + (f" ({seen.image})" if seen.image else "") + ".")

    def check_journal(self, arg: dict[str, Any]) -> tuple[bool, str]:
        entries = [e for e in self.world.journal.recent(200, kinds=[arg["kind"]])
                   if e["at"] >= self.mark.at
                   and (not arg.get("contains") or arg["contains"].lower() in e["text"].lower())
                   and (not arg.get("matches") or re.search(arg["matches"], e["text"], re.I))
                   and (not arg.get("url_contains")
                        or arg["url_contains"] in str((e.get("data") or {}).get("url", "")))]
        noun = "entry" if len(entries) == 1 else "entries"
        return bool(entries), (f"{len(entries)} {arg['kind']} {noun}"
                               + (f": {entries[-1]['text'][:160]}" if entries else "") + ".")

    def check_field(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """A row's field, found by its title: it reads as asked (A7: a page's button changed
        it through a turn)."""
        table = self._table(arg["collection"])
        rows = [r for r in self.world.collections.query(table, limit=None)
                if arg["title_contains"].lower() in " ".join(str(v) for v in r.values()
                                                              if isinstance(v, str)).lower()]
        if not rows:
            return False, f"No row of {table} like {arg['title_contains']!r}."
        value = str(rows[0].get(arg["field"]) or "")
        ok = bool(re.search(str(arg["matches"]), value, re.I))
        return ok, f"{arg['title_contains']}'s {arg['field']} is {value!r}."

    def check_hosts(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """A host a module's scripts may reach (A3): on its list and in its module.yaml."""
        m = self.world.modules.get(arg["module"])
        listed = grants.module_hosts(self.world, m["id"])
        self.world.files.sync_module(self.world.modules.get(m["id"]))
        in_file = arg["includes"] in (self.world.files.dir_of(m["id"]) / "module.yaml").read_text()
        ok = arg["includes"] in listed and in_file
        return ok, (f"{m['name']} may reach {', '.join(listed) or 'no host'}"
                    + ("; module.yaml says so" if in_file else "; module.yaml does not"))

    def check_file_in_folder(self, arg: dict[str, Any]) -> tuple[bool, str]:
        """A file Alpha wrote in a folder the person granted (A4): its name matches, and it
        holds the words named."""
        folder = Path(self.world.path).parent / "granted" / str(arg["folder"])
        found = [f for f in sorted(folder.glob("*")) if f.is_file()
                 and re.search(str(arg.get("name_matches") or "."), f.name)]
        if not found:
            return False, f"No file like /{arg.get('name_matches')}/ in {folder.name}."
        text = found[-1].read_text(errors="replace")
        missing = [w for w in arg.get("contains") or [] if w.lower() not in text.lower()]
        return not missing, (f"{found[-1].name} in {folder.name}"
                             + (f", missing {missing}" if missing else ", with every word named"))

def _days_since(weekday: str) -> int:
    """How many days back the last `weekday` before today was (1 to 7)."""
    days = ["monday", "tuesday", "wednesday", "thursday", "friday", "saturday", "sunday"]
    today = dt.now().astimezone().weekday()
    return (today - days.index(weekday.lower())) % 7 or 7


def _days_from_now(words: str, *, time: bool = False) -> str:
    """ "days:-2" → the date two days ago (a seed's dates stay relative to the run)."""
    when = dt.now(UTC) + timedelta(days=int(words.split(":", 1)[1]))
    return when.replace(microsecond=0).isoformat() if time else when.date().isoformat()


def text_pdf(pages: list[str]) -> bytes:
    """A plain PDF with these pages of text (one line each in Helvetica), for a journey's
    statement: written by hand, so no PDF library is needed to make one."""
    def esc(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    n = len(pages)
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        ("<< /Type /Pages /Kids [" + " ".join(f"{4 + 2 * i} 0 R" for i in range(n))
         + f"] /Count {n} >>").encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for i, text in enumerate(pages):
        lines = "".join(f"({esc(line)}) Tj 0 -16 Td " for line in text.split("\n"))
        stream = f"BT /F1 12 Tf 72 720 Td {lines}ET".encode()
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents"
                       f" {5 + 2 * i} 0 R /Resources << /Font << /F1 3 0 R >> >> >>".encode())
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream
                       + b"\nendstream")
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += (f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
            .encode())
    return bytes(out)


def run_journey(world: World, journey: dict[str, Any],
                runner: turn.Runner = route.run) -> Outcome:
    out = Outcome(name=journey["name"], title=journey.get("title", journey["name"]))
    began = time.monotonic()
    run = Run(world, runner)
    try:
        for spec in journey.get("steps", []):
            out.steps.append(run.step(spec))
        for spec in journey.get("checks", []):
            out.checks.append(run.check(spec))
        for spec in journey.get("steps_after", []):
            out.steps.append(run.step(spec))
        for spec in journey.get("checks_after", []):
            out.checks.append(run.check(spec))
    except Exception as e:  # a journey that breaks is a failed journey, not a crashed suite
        out.error = f"{type(e).__name__}: {e}"
    out.seconds = round(time.monotonic() - began, 1)
    return out


# ---- the suite ----


def default_world() -> Path:
    app = APP_HOME / "world.sqlite"
    if app.exists():
        return app
    return default_world_path()


def report(outcomes: list[Outcome], *, source: Path, home: Path, began: datetime,
           fresh: bool = False) -> str:
    passed = sum(1 for o in outcomes if o.passed)
    where = (f"a fresh world in `{home}` with the sign-ins of `{source.parent}`" if fresh
             else f"a copy of `{source}` in `{home}`")
    lines = [f"# Journeys, {began.strftime('%-d %b %Y %H:%M')}", "",
             f"{passed} of {len(outcomes)} passed. World: {where}."
             f" Thinks with: {route.WORDS[route.chosen(home / 'world.sqlite')]}, model"
             f" {config.model_for('turn')}.", ""]
    lines.append("| Journey | Verdict | Time |")
    lines.append("|---|---|---|")
    for o in outcomes:
        lines.append(f"| {o.title} | {'passed' if o.passed else 'FAILED'} | {o.seconds:.0f} s |")
    for o in outcomes:
        verdict = "passed" if o.passed else "failed"
        lines += ["", f"## {o.title} — {verdict} ({o.seconds:.0f} s)", ""]
        if o.error:
            lines.append(f"Broke: {o.error}")
        for s in o.steps:
            kind = next((k for k in ("say", "press", "check", "reader", "agent", "break_reader",
                                     "compact",
                                     "kit_export", "kit_install", "module_remove",
                                     "connect_folder", "due_now",
                                     "approve_action",
                                     "agree_workflow", "add_row", "seed") if k in s), "step")
            state = "ok" if s["ok"] else "not ok"
            head = f"- **{kind}** {s.get(kind, '')!s:.80} · {s.get('seconds', 0)} s · {state}"
            if "tokens_cache_read" in s:
                head += (f" · tokens: input {s['tokens_input']:,}, cache read"
                         f" {s['tokens_cache_read']:,}, cache made {s['tokens_cache_made']:,},"
                         f" output {s['tokens_out']:,}")
            if s.get("first_text_ms"):
                head += f" · first words {s['first_text_ms'] / 1000:.1f} s"
            detail = (s.get("reply") or s.get("result") or s.get("report") or s.get("detail")
                      or "")
            if s.get("problem"):
                detail = f"{detail} Problem: {s['problem']}"
            lines.append(head + (f"\n  {str(detail).strip()[:600]}" if detail else ""))
        for c in o.checks:
            lines.append(f"- check **{c['check']}**: {'✓' if c['ok'] else '✗'} {c['why']}"
                         f" ({c['seconds']} s)")
    return "\n".join(lines) + "\n"


def run_suite(names: list[str] | None = None, *, world_path: Path | None = None,
              scratch: Path | None = None, out_dir: Path | None = None,
              runner: turn.Runner = route.run, keep: bool = False,
              fresh: bool = False) -> tuple[int, Path]:
    """Copy the world (or start a fresh one with the app's sign-ins, `fresh`), run the
    journeys, write `docs/journeys/<stamp>.md` and `.json`. Returns (failures, report path)."""
    began = datetime.now().astimezone()
    source = world_path or default_world()
    stamp = began.strftime("%Y-%m-%d-%H%M")
    home = scratch or Path(os.environ.get("TMPDIR", "/tmp")) / "alpha-journeys" / stamp
    journeys = load(names)
    # A journey that says `fresh: true` (the first-day sentences) starts from an empty world;
    # the suite runs on a fresh world when told to or when every journey asked for one, else
    # on a copy, and a mixed set is refused rather than run half right.
    wants = {bool(j.get("fresh")) for j in journeys}
    if not fresh and wants == {True}:
        fresh = True
    if not fresh and True in wants:
        asked = ", ".join(j["name"] for j in journeys if j.get("fresh"))
        raise Problem(f"{asked} run on a fresh world (`--fresh`), the others on a copy: run"
                      " them apart.")
    copied = fresh_home(source, home) if fresh else copy_home(source, home)
    os.environ["ALPHA_HOME"] = str(home)  # the browser hand keeps its profiles under it
    os.environ["ALPHA_WORLD"] = str(copied)
    world = World(copied)
    outcomes: list[Outcome] = []
    try:
        for journey in journeys:
            outcomes.append(run_journey(world, journey, runner))
    finally:
        desk.close_all()  # the desk's client and any work still running with it
        world.close()
    folder = out_dir or (Path(__file__).resolve().parents[3] / "docs" / "journeys")
    folder.mkdir(parents=True, exist_ok=True)
    # The raw outcomes first: a slip in the report's wording must never lose a run's results.
    (folder / f"{stamp}.json").write_text(json.dumps(
        [o.__dict__ for o in outcomes], indent=1, ensure_ascii=False, default=str))
    text = report(outcomes, source=source, home=home, began=began, fresh=fresh)
    (folder / f"{stamp}.md").write_text(text)
    failed = sum(1 for o in outcomes if not o.passed)
    if not keep and not failed:
        # A failed run keeps its copy: the world and the workspaces are what explains it.
        shutil.rmtree(home, ignore_errors=True)
    return failed, folder / f"{stamp}.md"
