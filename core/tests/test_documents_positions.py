"""Documents with positions (P6 of provenance-documents-triggers-proposal.md, Q96; 19 Oct 2026):
a sheet read cell by cell with its formulas, a PDF by page, one file one document, a sheet's
cells by range, and rows imported by code with each value's cell as its source."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from conftest import building
from openpyxl import Workbook

from alpha.connectors.files import Files, parse_range, read_document
from alpha.mcp.tools import Tools
from alpha.runtime.agents import runner as agent_runner
from alpha.runtime.running import RunResult, TurnRequest
from alpha.world.store import Problem
from alpha.world.world import World


def workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "P&L"
    ws.append(["Account", "Amount", "Month"])
    ws.append(["Revenue", 48200, "Sept"])
    ws.append(["Cost of sales", 21400, "Sept"])
    ws["B5"] = "=B2-B3"
    notes = wb.create_sheet("Notes")
    notes["A1"] = "Prepared by RestoPros"
    wb.save(path)
    return path


def test_a_sheet_is_read_cell_by_cell_with_its_formulas(tmp_path: Path) -> None:
    text, structure = read_document(workbook(tmp_path / "pl.xlsx"))
    assert structure == {"kind": "sheets", "sheets": [
        {"name": "P&L", "rows": 5, "columns": 3}, {"name": "Notes", "rows": 1, "columns": 1}]}
    assert '# Sheet "P&L" (5 rows × 3 columns)' in text
    assert "2: A=Revenue | B=48200 | C=Sept" in text
    assert "5: B= (=B2-B3)" in text  # a formula Excel never computed shows as a formula
    assert parse_range("B2:F40") == (2, 2, 40, 6) and parse_range("B12") == (12, 2, 12, 2)
    assert parse_range("revenue") is None


def test_one_file_one_document_and_cells_by_range(world: World, tmp_path: Path) -> None:
    files = Files(world)
    source = workbook(tmp_path / "pl.xlsx")
    doc = files.take(source, module=None, origin="the person", by="person")
    again = files.take(source, module=None, origin="an email")
    assert again["id"] == doc["id"] and again["already"] is True
    assert len(world.store.all("SELECT id FROM documents")) == 1
    assert doc["structure"]["sheets"][0]["name"] == "P&L"
    read = files.read(doc["id"], sheet="P&L", cells="A2:B3")
    assert [c["cell"] for c in read["cells"]] == ["A2", "B2", "A3", "B3"]
    assert "B2 = 48200" in read["text"]
    with pytest.raises(Problem, match="has no sheet 'Profit'; its sheets are P&L, Notes"):
        files.read(doc["id"], sheet="Profit")


def test_a_pdf_is_read_page_by_page(world: World) -> None:
    entity = world.entities.resolve("document", "Statement", {"path": "/x.pdf"})["entity"]["id"]
    with world.store.tx() as db:
        db.execute(
            "INSERT INTO documents (id, entity_id, connection, path, title, kind, text, size,"
            " modified_at, indexed_at, structure) VALUES ('d_1', ?, NULL, '/x.pdf', 'Statement"
            " Sept.pdf', 'pdf', ?, 10, '2026', '2026', ?)",
            (entity,
             "[page 1]\nOpening balance 1,200\n\n[page 2]\nClosing balance 3,400",
             '{"kind": "pdf", "pages": 2}'))
    page = Files(world).read("d_1", page=2)
    assert page["text"].strip() == "Closing balance 3,400" and page["more"] is False
    with pytest.raises(Problem, match="has pages 1 to 2"):
        Files(world).read("d_1", page=3)


def ledger(world: World) -> str:
    building(world).collection_create("ledger", "Ledger", [
        {"name": "account", "kind": "text"}, {"name": "amount", "kind": "number"},
        {"name": "month", "kind": "text"}])
    return "ledger"


def test_rows_are_imported_by_code_with_each_values_cell(world: World, tmp_path: Path) -> None:
    table = ledger(world)
    doc = Files(world).take(workbook(tmp_path / "pl.xlsx"), module=None, origin="the person",
                            by="person")
    out = Tools(world, turn="j_1").world_file("import", ref=doc["id"], into=table,
                                              sheet="P&L")
    # the formula-only row (B5, never computed by Excel) has no value to bring
    assert isinstance(out, dict) and out["imported"] == 2, out
    assert out["mapped"] == {"Account": "account", "Amount": "amount", "Month": "month"}
    rows = sorted(world.collections.query(table, limit=None), key=lambda r: str(r["account"]))
    revenue = next(r for r in rows if r["account"] == "Revenue")
    assert revenue["amount"] == 48200
    cell = revenue["_provenance"]["fields"]["amount"]
    assert {k: cell[k] for k in ("kind", "document", "sheet", "cell")} == {
        "kind": "document", "document": doc["id"], "sheet": "P&L", "cell": "B2"}
    assert cell["source"] == f"{doc['title']} · P&L!B2"  # the words the window and runs read
    # columns named by the person, and a key to update by
    out = Tools(world, turn="j_2").world_file("import", ref=doc["id"], into=table,
                                              sheet="P&L", columns={"Account": "account",
                                                                    "Amount": "amount"},
                                              key="account")
    assert isinstance(out, dict) and out["unmapped"] == ["Month"]
    assert len(world.collections.query(table, limit=None)) == 2  # updated, not added


def test_an_agents_import_into_a_staged_table_waits_with_its_cells(world: World,
                                                                    tmp_path: Path) -> None:
    table = ledger(world)
    doc = Files(world).take(workbook(tmp_path / "pl.xlsx"), module=None, origin="email",
                            by="person")
    agent = building(world, turn="j_1").agent_create(
        "When run, bring the month's P&L into the ledger", "", "Import the P&L sheet.",
        name="Books", stage=[table])

    def runner(req: TurnRequest) -> RunResult:
        out: Any = Tools(world, turn=req.turn_id, thread=req.thread_id, actor=req.actor,
                         run=req.run).world_file("import", ref=doc["id"], into=table,
                                                 sheet="P&L")
        assert out.get("staged"), out
        return RunResult(ok=True, reply="Brought in the P&L; it waits for your yes.")

    agent_runner.run(world, agent["id"], runner=runner)
    assert world.collections.query(table, limit=None) == []
    [held] = [s for s in world.store.all("SELECT id FROM staged WHERE state = 'held'")]
    landed = world.staged.land_rows(held["id"])
    assert len(landed["landed"]) == 2
    row = next(r for r in world.collections.query(table, limit=None) if r["account"] == "Revenue")
    assert row["_provenance"]["fields"]["amount"]["cell"] == "B2"


def test_an_exports_header_row_is_found_under_its_title_rows(world: World,
                                                             tmp_path: Path) -> None:
    """A report as accounting software exports it (found on a RestoPros QuickBooks export,
    19 Oct): title rows of one cell, then the header with an empty corner over the accounts."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    for row in (["RestoPros - St. Louis, MO"], ["Profit and Loss by Month"],
                ["June 1, 2025-June 30, 2026"], [], [None, "Sep 2025", "Oct 2025"],
                ["Income"], ["Mitigation Sales", 190595.68, 135855.69]):
        ws.append(row)
    path = tmp_path / "pl-by-month.xlsx"
    wb.save(path)
    building(world).collection_create("pl", "P&L", [
        {"name": "account", "kind": "text"}, {"name": "sep_2025", "kind": "number"}])
    doc = Files(world).take(path, module=None, origin="the person", by="person")
    out = Tools(world, turn="j_1").world_file("import", ref=doc["id"], into="pl",
                                              columns={"A": "account", "Sep 2025": "sep_2025"})
    assert isinstance(out, dict) and out["header_row"] == 5, out
    assert out["mapped"] == {"A": "account", "Sep 2025": "sep_2025"}
    [sales] = [r for r in world.collections.query("pl", limit=None)
               if r["account"] == "Mitigation Sales"]
    assert sales["sep_2025"] == 190595.68
    assert sales["_provenance"]["fields"]["sep_2025"]["cell"] == "B7"
