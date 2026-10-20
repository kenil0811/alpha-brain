from __future__ import annotations

import io
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any

import pytest

from alpha.connectors.base import Connections, manifests
from alpha.connectors.browser import Browser
from alpha.connectors.calendar import Attendee, Calendar, CalendarEvent
from alpha.connectors.files import Files
from alpha.context import layers
from alpha.mcp.tools import Tools
from alpha.world.sites import site_of
from alpha.world.store import Problem
from alpha.world.world import World

# ---- manifests ----


def test_every_builtin_connector_has_a_manifest_and_a_skill() -> None:
    names = {m["name"] for m in manifests()}
    assert {"files", "browser", "calendar"} <= names
    for m in manifests():
        assert m["tools"] and all(t["effect"] in {"read", "write"} for t in m["tools"])
        assert (Path(__file__).parents[2] / "connectors" / m["name"] / "SKILL.md").exists()


# ---- files ----


def test_watch_sync_change_remove(world: World, tmp_path: Path) -> None:
    folder = tmp_path / "Job search"
    (folder / "drafts").mkdir(parents=True)
    (folder / "notes.md").write_text("Lumen Analytics: backend role, Kafka and Go.")
    (folder / "drafts" / "cover.txt").write_text("Dear hiring manager at Northwind")
    (folder / ".hidden.md").write_text("secret")
    (folder / "image.png").write_bytes(b"\x89PNG")
    files = Files(world)
    files.watch(str(folder))
    first = files.sync()
    assert first["added"] == 2 and first["folders"] == 1
    assert files.search("kafka")[0]["title"] == "notes.md"
    assert files.sync()["added"] == 0  # nothing changed
    import os
    (folder / "notes.md").write_text("Lumen Analytics: interview on Wednesday.")
    later = datetime.now().timestamp() + 5
    os.utime(folder / "notes.md", (later, later))
    assert files.sync()["changed"] == 1
    assert files.search("wednesday") and not files.search("kafka")
    (folder / "drafts" / "cover.txt").unlink()
    assert files.sync()["removed"] == 1
    doc = files.read("notes.md")
    assert "Wednesday" in doc["text"] and doc["more"] is False
    entity = world.entities.find(keys={"path": str(folder / "notes.md")})[0]
    assert entity["kind"] == "document"
    assert any("Read notes.md" in e["text"] for e in world.journal.recent(20))


def test_office_documents_are_read(world: World, tmp_path: Path) -> None:
    import docx
    from openpyxl import Workbook

    folder = tmp_path / "docs"
    folder.mkdir()
    d = docx.Document()
    d.add_paragraph("Resume: senior backend engineer, Postgres")
    d.save(str(folder / "resume.docx"))
    wb = Workbook()
    sheet: Any = wb.active
    sheet.append(["company", "role"])
    sheet.append(["Harbourline", "Backend Engineer (Go)"])
    wb.save(str(folder / "targets.xlsx"))
    files = Files(world)
    files.watch(str(folder))
    assert files.sync()["added"] == 2
    assert files.search("postgres")[0]["title"] == "resume.docx"
    assert "Harbourline" in files.read("targets.xlsx")["text"]


def test_two_syncs_of_one_folder_at_once(world: World, tmp_path: Path,
                                         monkeypatch: pytest.MonkeyPatch) -> None:
    """A second sync of a folder starting while the first still walks it: the watcher's timer
    starts one 2 s after the last change, and a walk of his ~/Desktop/dev takes 8 s. Both read
    the folder's documents before walking, so both saw a new file as new, and the second's
    INSERT failed on UNIQUE(path): the whole folder's sync failed and the connection showed
    broken (his core log, 20 Oct 2026). Here the second sync runs inside the first's walk."""
    import alpha.connectors.files as files_module

    folder = tmp_path / "dev"
    folder.mkdir()
    (folder / "STATE.md").write_text("The state, as of the morning.")
    files = Files(world)
    conn = files.watch(str(folder))
    assert files.sync()["added"] == 1
    (folder / "2026-10-20-1307.md").write_text("A journey's report.")
    (folder / "STATE.md").unlink()

    walk = files_module._walk
    other: dict[str, Any] = {}

    def walk_while_another_syncs(root: Path) -> list[Path]:
        found = walk(root)
        if not other:
            other["started"] = True
            other["sync"] = files.sync()
        return found

    monkeypatch.setattr(files_module, "_walk", walk_while_another_syncs)
    first = files.sync()
    assert other["sync"]["added"] == 1 and other["sync"]["removed"] == 1
    assert first["folders"] == 1  # the first sync finished too
    assert first["added"] == first["removed"] == 0  # the other sync did both already
    assert Connections(world.store).get(conn["id"])["status"] == "connected"
    assert len(world.store.all("SELECT path FROM documents")) == 2
    said = [e["text"] for e in world.journal.recent(20)]
    assert said.count("Read 2026-10-20-1307.md in dev (3 words).") == 1
    assert sum("STATE.md was removed" in t for t in said) == 1


def test_a_file_another_folder_reads_stays_with_it(world: World, tmp_path: Path,
                                                   monkeypatch: pytest.MonkeyPatch) -> None:
    """A folder watched inside one already watched: its files stay the inner folder's
    documents (one file, one document; Kenil, 20 Oct 2026), even one the inner folder reads
    while the outer one walks. Before, the outer sync failed on UNIQUE(path)."""
    import alpha.connectors.files as files_module

    outer = tmp_path / "dev"
    inner = outer / "alpha" / "docs"
    inner.mkdir(parents=True)
    (outer / "README.md").write_text("The dev folder.")
    (inner / "STATE.md").write_text("Alpha's state.")
    files = Files(world)
    inner_conn = files.watch(str(inner))
    assert files.sync(str(inner))["added"] == 1
    outer_conn = files.watch(str(outer))
    assert files.sync(str(outer))["added"] == 1  # README.md alone
    (inner / "log.md").write_text("A new entry.")

    walk = files_module._walk
    other: dict[str, Any] = {}

    def walk_while_the_inner_folder_syncs(root: Path) -> list[Path]:
        found = walk(root)
        if root == outer and not other:
            other["sync"] = files.sync(str(inner))
        return found

    monkeypatch.setattr(files_module, "_walk", walk_while_the_inner_folder_syncs)
    result = files.sync(str(outer))
    assert other["sync"]["added"] == 1
    assert result["folders"] == 1 and result["added"] == result["changed"] == 0
    (inner / "STATE.md").unlink()
    assert files.sync(str(outer))["removed"] == 0  # the inner folder's to mark
    assert files.sync(str(inner))["removed"] == 1
    held = {r["path"]: r["connection"] for r in world.store.all(
        "SELECT path, connection FROM documents")}
    assert held == {str(outer / "README.md"): outer_conn["id"],
                    str(inner / "STATE.md"): inner_conn["id"],
                    str(inner / "log.md"): inner_conn["id"]}
    for conn in (outer_conn, inner_conn):
        assert Connections(world.store).get(conn["id"])["status"] == "connected"


def test_a_skipped_folder_is_passed_by_not_entered(world: World, tmp_path: Path,
                                                   monkeypatch: pytest.MonkeyPatch) -> None:
    """His ~/Desktop/dev: each sync went into node_modules, .venv and .git only to drop what
    it found there, 9.4 s a sync (20 Oct 2026). A skipped folder is now never listed."""
    import os

    import alpha.connectors.files as files_module

    folder = tmp_path / "dev"
    for sub in ("node_modules/pkg", ".venv/lib", ".git/objects", "app/.data/browser", "app/docs"):
        (folder / sub).mkdir(parents=True)
        (folder / sub / "README.md").write_text("inside " + sub)
    entered: list[str] = []
    walk = os.walk

    def listing(top: Any, *args: Any, **kwargs: Any) -> Any:
        for here, dirs, names in walk(top, *args, **kwargs):
            entered.append(os.path.relpath(here, folder))
            yield here, dirs, names

    monkeypatch.setattr(os, "walk", listing)
    found = files_module._walk(folder)
    assert [p.relative_to(folder).as_posix() for p in found] == ["app/docs/README.md"]
    assert sorted(entered) == [".", "app", "app/docs"]


def test_a_change_no_sync_would_see_starts_none() -> None:
    from alpha.connectors.files import _seen_by_sync

    root = "/Users/someone/dev"
    assert _seen_by_sync(root, f"{root}/alpha/docs/STATE.md")
    assert _seen_by_sync(root, f"{root}/notes.txt")
    for path in (f"{root}/alpha/node_modules/pkg/README.md",  # a skipped folder
                 f"{root}/alpha-v2/.data/browser/Preferences.json",  # a hidden one
                 f"{root}/alpha/world.sqlite-wal",  # a file no sync reads
                 f"{root}/alpha/.DS_Store", root):
        assert not _seen_by_sync(root, path), path
    assert _seen_by_sync(root, "/elsewhere/odd.md")  # can't place it: the sync judges


def test_a_change_during_a_sync_makes_one_more_after_it() -> None:
    """Overlapping syncs of one folder kept his core at 200% CPU (20 Oct 2026): while one
    runs, any number of changes make exactly one more after it, never one alongside."""
    import threading

    from alpha.connectors.files import Syncs

    started, release = threading.Event(), threading.Event()
    calls: list[str] = []
    now: dict[str, int] = {}
    most: dict[str, int] = {}

    def sync(folder: str) -> None:
        now[folder] = now.get(folder, 0) + 1
        most[folder] = max(most.get(folder, 0), now[folder])
        calls.append(folder)
        if len(calls) == 1:
            started.set()
            assert release.wait(5)
        now[folder] -= 1
        if len(calls) == 2:
            raise RuntimeError("a failed sync still lets the next one run")

    syncs = Syncs(sync)
    first = threading.Thread(target=syncs.run, args=("/dev",))
    first.start()
    assert started.wait(5)
    for _ in range(3):
        syncs.run("/dev")  # returns at once: noted for after
    syncs.run("/other")  # another folder is its own
    release.set()
    first.join(5)
    assert calls == ["/dev", "/other", "/dev"] and most["/dev"] == 1
    assert not syncs.running and not syncs.again
    syncs.run("/dev")
    assert calls[-1] == "/dev" and len(calls) == 4


def test_watch_refuses_home_and_missing(world: World) -> None:
    with pytest.raises(Problem, match="home folder"):
        Files(world).watch("~")
    with pytest.raises(Problem, match="isn't a folder"):
        Files(world).watch("/no/such/folder")


# ---- browser ----


def test_site_of() -> None:
    assert site_of("https://www.linkedin.com/in/priya/") == "linkedin.com"
    assert site_of("jobs.example.co.uk") == "example.co.uk"


def test_read_uses_the_signin_and_journals(world: World) -> None:
    jobs: list[dict[str, Any]] = []

    def runner(job: dict[str, Any], timeout: int) -> dict[str, Any]:
        jobs.append(job)
        if job["op"] == "status":
            return {"signed_in": True}
        return {"status": 200, "final_url": job["url"], "title": "My Network",
                "text": "Priya Raman Engineering Manager", "links": [
                    {"text": "Priya Raman", "url": "https://www.linkedin.com/in/priya/",
                     "near": "Priya Raman · Engineering Manager at Lumen"}]}

    Connections(world.store).upsert("browser", "linkedin.com", status="needs_ok")
    page = Browser(world, runner).read("https://www.linkedin.com/mynetwork/")
    assert jobs[0]["op"] == "status" and "profile" in jobs[1]
    assert page["signed_in"] and page["links"][0]["near"].startswith("Priya")
    conn = Connections(world.store).find("browser", "linkedin.com")
    assert conn is not None and conn["status"] == "connected"
    assert "signed in" in world.journal.recent(1)[0]["text"]


def test_a_signin_ask_closes_when_asked_again_and_when_signed_in(
    world: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    class Window:
        def __init__(self, *args: Any, **kwargs: Any) -> None:
            self.stdin = io.StringIO()

    monkeypatch.setenv("ALPHA_NODE", "node")
    monkeypatch.setattr("alpha.connectors.browser.subprocess.Popen", Window)
    browser = Browser(world, lambda job, timeout: {"signed_in": True})
    browser.start_signin("gmail.com")
    browser.start_signin("gmail.com")
    asks = world.journal.open_asks()
    assert len(asks) == 1 and "gmail.com" in asks[0]["text"]
    browser.refresh("gmail.com")
    assert world.journal.open_asks() == []


def test_a_signin_covers_the_sites_its_window_passed_through(
    world: World, tmp_path: Path
) -> None:
    profile = tmp_path / "gmail.com"
    profile.mkdir()
    (profile / "alpha-signin.json").write_text(
        '{"hosts": ["www.gmail.com", "accounts.google.com", "mail.google.com"]}')
    Connections(world.store).upsert("browser", "gmail.com", status="needs_ok",
                                    config={"profile": str(profile)})
    jobs: list[dict[str, Any]] = []

    def runner(job: dict[str, Any], timeout: int) -> dict[str, Any]:
        jobs.append(job)
        if job["op"] == "status":
            return {"signed_in": "google.com" in job["sites"]}
        return {"status": 200, "final_url": job["url"], "title": "Inbox", "text": "", "links": []}

    page = Browser(world, runner).read("https://mail.google.com/mail/u/0/")
    assert page["signed_in"] and jobs[1]["profile"] == str(profile)
    assert jobs[0]["sites"] == ["gmail.com", "google.com"]
    conn = Connections(world.store).find("browser", "gmail.com")
    assert conn is not None and conn["status"] == "connected"


def test_a_page_asking_for_a_signin_tries_the_signins_alpha_holds_first(
    world: World, tmp_path: Path
) -> None:
    for site in ("example.com", "gmail.com"):  # tried in this order
        (tmp_path / site).mkdir()
        Connections(world.store).upsert("browser", site, config={"profile": str(tmp_path / site)})
    jobs: list[dict[str, Any]] = []

    def runner(job: dict[str, Any], timeout: int) -> dict[str, Any]:
        jobs.append(job)
        in_gmail = str(job.get("profile", "")).endswith("gmail.com")
        if job["op"] == "status":
            return {"signed_in": in_gmail}
        return {"status": 200, "final_url": job["url"], "title": "Inbox", "text": "",
                "links": [], "blocked": not in_gmail}

    browser = Browser(world, runner)
    page = browser.read("https://mail.google.com/mail/u/0/")
    assert page["signed_in"] and not page["needs_signin"]
    assert [j["op"] for j in jobs] == ["read", "status", "status", "read"]
    # From now on the Gmail sign-in covers google.com: no more trying.
    jobs.clear()
    browser.read("https://mail.google.com/mail/u/0/")
    assert [j["op"] for j in jobs] == ["read"] and jobs[0]["profile"].endswith("gmail.com")


def test_read_refuses_local_addresses(world: World) -> None:
    with pytest.raises(Problem, match="local network"):
        Browser(world, lambda j, t: {}).read("http://192.168.1.1/admin")


# ---- calendar ----


class FakeCalendar:
    def __init__(self, granted: bool = True) -> None:
        self.granted = granted
        self.asked = False
        start = datetime.now(UTC).replace(minute=0, second=0, microsecond=0) + timedelta(hours=1)
        self.items = [
            CalendarEvent(uid="ev-1", title="Lumen intro call", starts_at=start,
                          ends_at=start + timedelta(minutes=30), calendar="Work",
                          organiser=Attendee("Priya Raman", "priya@lumen.example"),
                          attendees=[Attendee("Priya Raman", "Priya@Lumen.example"),
                                     Attendee(None, None)]),
            CalendarEvent(uid="ev-2", title="Dentist", starts_at=start + timedelta(days=2),
                          ends_at=start + timedelta(days=2, hours=1)),
        ]

    def access(self) -> str:
        return "granted" if self.granted else ("not_asked" if not self.asked else "denied")

    def request_access(self) -> str:
        self.asked = True
        return self.access()

    def events(self, start: datetime, end: datetime) -> list[CalendarEvent]:
        return list(self.items)


def test_calendar_connect_sync_and_people(world: World) -> None:
    source = FakeCalendar()
    cal = Calendar(world, source)
    conn = cal.connect()
    assert conn["status"] == "connected" and conn["last_sync"]
    priya = world.entities.find(keys={"email": "priya@lumen.example"})
    assert len(priya) == 1  # organiser and attendee are one person
    now = datetime.now(UTC)
    week = cal.between(now.isoformat(), (now + timedelta(days=7)).isoformat())
    assert [e["title"] for e in week] == ["Lumen intro call", "Dentist"]
    assert week[0]["attendees"][0]["entity_id"] == priya[0]["id"]
    assert cal.sync()["unchanged"] == 2
    source.items[1].title = "Dentist (moved)"
    source.items = source.items[1:]
    result = cal.sync()
    assert result["changed"] == 1 and result["removed"] == 1


class FlakyCalendar(FakeCalendar):
    """A calendar whose next reads come back as planned (an empty one), then as it stands."""

    def __init__(self) -> None:
        super().__init__()
        self.reads: list[list[CalendarEvent]] = []

    def events(self, start: datetime, end: datetime) -> list[CalendarEvent]:
        return self.reads.pop(0) if self.reads else list(self.items)


def test_a_read_that_comes_back_empty_removes_nothing(world: World) -> None:
    """19 Oct 2026: one read with no events while access was granted took all 28 as deleted.
    A missing event is gone only when a second read, on a fresh store, misses it too."""
    source = FlakyCalendar()
    cal = Calendar(world, source)
    cal.connect()
    source.reads = [[]]  # one empty read, then the calendar as it stands
    assert cal.sync()["removed"] == 0
    assert [e["title"] for e in cal.between("2000-01-01", "2999-01-01")] == [
        "Lumen intro call", "Dentist"]
    # deleted for real: both reads miss it, so it goes
    source.items = source.items[:1]
    assert cal.sync()["removed"] == 1
    # back on the calendar: the next read finds it, and the journal says it is back
    source.items = FakeCalendar().items
    assert cal.sync()["changed"] == 1
    back = world.journal.recent(1, kinds=["saw"])[0]["text"]
    assert back.startswith("Back on your calendar: Dentist")


def test_calendar_without_access_needs_ok(world: World) -> None:
    cal = Calendar(world, FakeCalendar(granted=False))
    assert cal.connect()["status"] == "needs_ok"
    with pytest.raises(Problem, match="System Settings"):
        cal.sync()


# ---- the pre-pack and tools see it ----


def test_the_layers_show_reach_calendar_and_documents(world: World, tmp_path: Path) -> None:
    assert "No folder is granted beyond your own directory" in layers.given(world, "x")
    Calendar(world, FakeCalendar()).connect()
    folder = tmp_path / "notes"
    folder.mkdir()
    (folder / "lumen.md").write_text("Lumen salary band 90 to 100k")
    Files(world).watch(str(folder))
    Files(world).sync()
    text = layers.given(world, "what is the lumen salary band")
    assert "calendar: macos — connected" in text and "files:" in text
    assert "document" in text and "lumen.md" in text
    if datetime.now().astimezone().hour < 23:
        assert "TODAY'S CALENDAR" in text and "Lumen intro call with Priya Raman" in text


def test_tools_cover_the_connectors(world: World, tmp_path: Path) -> None:
    t = Tools(world)
    folder = tmp_path / "f"
    folder.mkdir()
    (folder / "a.txt").write_text("quarterly numbers for Harbourline")
    watched = t.folder_watch(str(folder))
    assert watched["added"] == 1
    assert t.search("harbourline")["documents"][0]["title"] == "a.txt"
    assert t.document_read("a.txt")["text"].startswith("quarterly")
    assert t.connections_list()[0]["connector"] == "files"
    assert "error" in t.page_read("file:///etc/passwd")
